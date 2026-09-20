from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


PSI_PATH = Path("/Users/iant1359/Downloads/PSI 23.07 _KT check.xlsx")
MISA_PATH = Path(
    "/Users/iant1359/Downloads/mismatch/So_chi_tiet_ban_hang 22.07.xlsx"
)


def emit(label: str, value: object) -> None:
    print(f"{label}\t{json.dumps(value, ensure_ascii=False, default=str)}")


psi = openpyxl.load_workbook(
    PSI_PATH, data_only=True, read_only=True, keep_links=False
)
revenue = psi["Revenue final"]

revenue_header = None
flagged_rows: list[dict[str, object]] = []
for row_number, row in enumerate(
    revenue.iter_rows(min_row=6, max_col=17, values_only=True), start=6
):
    if row_number == 6:
        revenue_header = list(row)
        continue

    variance = row[16]
    if isinstance(variance, (int, float)) and variance != 0:
        flagged_rows.append(
            {
                "row": row_number,
                "sales_date": row[1],
                "sku": row[2],
                "product_name": row[3],
                "quantity": row[8],
                "net_rev_sold": row[9],
                "sale_order": row[11],
                "expected_by_accounting": row[15],
                "variance": variance,
            }
        )

emit("REVENUE_HEADER", revenue_header)
emit("FLAGGED_COUNT", len(flagged_rows))
emit("FLAGGED_TOTAL", sum(float(row["variance"]) for row in flagged_rows))
for row in flagged_rows:
    emit("FLAGGED", row)

target_skus = {str(row["sku"]).strip() for row in flagged_rows}
target_date_skus = {
    (str(row["sales_date"])[:10], str(row["sku"]).strip()) for row in flagged_rows
}

misa = openpyxl.load_workbook(
    MISA_PATH, data_only=True, read_only=True, keep_links=False
)
misa_sheet = misa[misa.sheetnames[0]]

misa_header = None
source_candidates: list[dict[str, object]] = []
for row_number, row in enumerate(
    misa_sheet.iter_rows(min_row=1, max_col=37, values_only=True), start=1
):
    if row_number == 4:
        misa_header = list(row)
        emit(
            "MISA_HEADER",
            [
                {
                    "column": get_column_letter(index),
                    "index": index - 1,
                    "name": value,
                }
                for index, value in enumerate(row, start=1)
            ],
        )
        continue
    if row_number <= 4:
        continue

    sku = str(row[10]).strip() if row[10] not in (None, "") else ""
    accounting_date = str(row[0])[:10] if row[0] not in (None, "") else ""
    if (accounting_date, sku) not in target_date_skus:
        continue

    sales = float(row[16] or 0)
    discount = float(row[19] or 0)
    returned = float(row[22] or 0)
    price_reduction = float(row[23] or 0)
    source_candidates.append(
        {
            "row": row_number,
            "accounting_date": row[0],
            "document_date": row[1],
            "document_number": row[2],
            "invoice_number": row[7],
            "description": row[8],
            "sku": row[10],
            "quantity_sold": row[13],
            "sales": sales,
            "discount": discount,
            "discount_rate": row[20],
            "returned_quantity": row[21],
            "returned_value": returned,
            "price_reduction": price_reduction,
            "sale_order": row[28],
            "net_with_all_line_fields": sales
            - discount
            - returned
            - price_reduction,
        }
    )

emit("SOURCE_CANDIDATE_COUNT", len(source_candidates))
for row in source_candidates:
    emit("SOURCE_CANDIDATE", row)
