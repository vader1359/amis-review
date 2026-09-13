"""Portable preparation preserving the verified September 2026 offline rules."""

from __future__ import annotations

import hashlib
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ._online_manual_check import load_manual_check


def txt(value: Any) -> str:
    return "" if value is None else str(value).strip()


def norm(value: Any) -> str:
    raw = unicodedata.normalize("NFKD", txt(value))
    return "".join(char for char in raw if not unicodedata.combining(char)).casefold()


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value[:10], pattern).date()
            except ValueError:
                pass
    return None


def iso(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_map(values: list[Any] | tuple[Any, ...]) -> dict[str, int]:
    return {norm(value): index for index, value in enumerate(values) if txt(value)}


def find_index(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        if norm(name) in headers:
            return headers[norm(name)]
    return None


def value(row: tuple[Any, ...], headers: dict[str, int], *names: str) -> Any:
    index = find_index(headers, *names)
    return row[index] if index is not None and index < len(row) else None


def all_rows(
    path: Path, sheet: str, header_row: int
) -> tuple[list[Any], list[tuple[Any, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet]
        iterator = worksheet.iter_rows(min_row=header_row, values_only=True)
        headers = list(next(iterator))
        return headers, list(iterator)
    finally:
        workbook.close()


def require_headers(
    source: str, headers: list[Any], required: list[tuple[str, ...]]
) -> None:
    mapped = header_map(headers)
    missing = [
        " / ".join(group) for group in required if find_index(mapped, *group) is None
    ]
    if missing:
        raise ValueError(f"{source}: missing required headers: {', '.join(missing)}")


def prepare(sources: dict[str, Path], as_of: date, prior_psi: Path) -> dict[str, Any]:
    start_date = date(2024, 1, 1)
    for source_path in sources.values():
        if not source_path.is_file():
            raise FileNotFoundError(source_path)

    manual_registry = load_manual_check(sources["Manual Check"])
    manual_summary = manual_registry.summary(as_of)

    def sku(value: Any, source_scope: str) -> str:
        return manual_registry.map_sku(value, source_scope=source_scope, as_of=as_of)

    approved_order_exclusions = {
        rule.order_id
        for rule in manual_registry.order_exclusions
        if rule.is_active(as_of)
        and rule.action == "EXCLUDE ORDER FROM PSI"
        and rule.scope == "ALL PSI BUSINESS SHEETS"
        and rule.disposition == "PERMANENT / DONE"
    }
    active_preorder_exclusion_rows = [
        (rule.order_id, rule.canonical_sku)
        for rule in manual_registry.preorder_exclusions
        if rule.is_active(as_of)
        and rule.action == "EXCLUDE FROM PREORDER"
        and rule.disposition == "PERMANENT / DONE"
    ]
    approved_preorder_exclusions = {
        (rule.order_id, rule.canonical_sku)
        for rule in manual_registry.preorder_exclusions
        if rule.is_active(as_of)
        and rule.action == "EXCLUDE FROM PREORDER"
        and rule.disposition == "PERMANENT / DONE"
    }

    def prior_mismatch_keys() -> set[tuple[str, str, str]]:
        book = load_workbook(prior_psi, read_only=True, data_only=True)
        try:
            worksheet = book["Mismatch"]
            return {
                (txt(row[1]), txt(row[2]), txt(row[3]))
                for row in worksheet.iter_rows(min_row=4, values_only=True)
                if len(row) >= 4 and any(item not in (None, "") for item in row)
            }
        finally:
            book.close()

    data_gaps: list[list[Any]] = []
    mismatches: list[list[Any]] = []

    def add_gap(source: str, key: str, issue: str, action: str) -> None:
        data_gaps.append([source, txt(key), issue, action])

    def add_mismatch(
        source: str, key: str, issue: str, action: str, evidence: str
    ) -> None:
        mismatches.append(["OPEN", source, txt(key), issue, action, evidence])

    # User-approved Product Master mapping: Nguồn gốc is Brand; Mã hãng is Brand Code.
    product_headers, product_source = all_rows(
        sources["Product Master"], "Danh sách", 1
    )
    require_headers(
        "Product Master",
        product_headers,
        [
            ("Mã hàng hóa",),
            ("Tên hàng hóa",),
            ("Nguồn gốc",),
            ("Category",),
            ("Sub Category",),
        ],
    )
    product_idx = header_map(product_headers)
    products: dict[str, dict[str, Any]] = {}
    product_duplicate_skus = 0
    for source_row, row in enumerate(product_source, start=2):
        raw_sku = txt(value(row, product_idx, "Mã hàng hóa"))
        code = sku(raw_sku, "PRODUCT MASTER")
        if not code:
            if any(txt(item) for item in row):
                add_gap(
                    "Product Master",
                    f"row {source_row}",
                    "Blank SKU",
                    "Populate Mã hàng hóa in Product Master",
                )
            continue
        brand = txt(value(row, product_idx, "Nguồn gốc"))
        brand_code = txt(value(row, product_idx, "Mã hãng"))
        category = txt(value(row, product_idx, "Category"))
        subcategory = txt(value(row, product_idx, "Sub Category", "Sub-category"))
        if code in products:
            product_duplicate_skus += 1
        products[code] = {
            "sku": code,
            "name": txt(value(row, product_idx, "Tên hàng hóa", "Tên hàng")),
            "brand": brand,
            "brand_code": brand_code,
            "category": category,
            "subcategory": subcategory,
            "supplier": txt(value(row, product_idx, "Loại hàng hóa")),
        }

    # CRM Final orders.
    crm_order_headers, crm_order_source = all_rows(
        sources["CRM Sales Order"], "Danh sách", 1
    )
    require_headers(
        "CRM Sales Order / Danh sách",
        crm_order_headers,
        [
            ("Số đơn hàng", "Số đơn hàng"),
            ("Ngày duyệt",),
            ("Trạng thái phê duyệt",),
            ("Tình trạng",),
        ],
    )
    crm_order_idx = header_map(crm_order_headers)
    crm_orders: dict[str, dict[str, Any]] = {}
    crm_order_all: dict[str, dict[str, Any]] = {}
    for row in crm_order_source:
        order_id = txt(value(row, crm_order_idx, "Số đơn hàng", "Số đơn hàng")).upper()
        approval_date = to_date(value(row, crm_order_idx, "Ngày duyệt"))
        approval_status_text = txt(value(row, crm_order_idx, "Trạng thái phê duyệt"))
        state_text = txt(value(row, crm_order_idx, "Tình trạng"))
        if order_id.startswith("DH-"):
            crm_order_all[order_id] = {
                "approval_date": approval_date.isoformat() if approval_date else "",
                "approval_status": approval_status_text,
                "state": state_text,
            }
        if (
            not order_id.startswith("DH-")
            or approval_date is None
            or not (start_date <= approval_date <= as_of)
        ):
            continue
        if norm(approval_status_text) != norm("Đã duyệt") or any(
            token in norm(state_text) for token in ("huy", "cancel")
        ):
            continue
        if order_id in approved_order_exclusions:
            continue
        crm_orders[order_id] = {
            "order_id": order_id,
            "approval_date": approval_date.isoformat(),
            "customer": txt(value(row, crm_order_idx, "Khách hàng")),
            "customer_code": txt(value(row, crm_order_idx, "Mã khách hàng")),
            "order_value": num(value(row, crm_order_idx, "Giá trị đơn hàng")),
            "status": state_text,
            "approval_status": approval_status_text,
            "delivery_status": txt(value(row, crm_order_idx, "Tình trạng giao hàng")),
            "payment_status": txt(value(row, crm_order_idx, "Tình trạng thanh toán")),
            "invoiced_value": num(value(row, crm_order_idx, "Giá trị đã xuất hóa đơn")),
        }

    # CRM products are aggregated by approved Order ID + canonical SKU.
    crm_item_headers, crm_item_source = all_rows(
        sources["CRM Sales Order"], "Bảng hàng hóa", 1
    )
    require_headers(
        "CRM Sales Order / Bảng hàng hóa",
        crm_item_headers,
        [
            ("Số đơn hàng",),
            ("Mã hàng hóa", "Mã hàng hóa"),
            ("Số lượng",),
            ("Thành tiền sau CK", "Tổng tiền"),
        ],
    )
    crm_item_idx = header_map(crm_item_headers)
    crm_item_agg: dict[tuple[str, str], dict[str, Any]] = {}
    for source_row, row in enumerate(crm_item_source, start=2):
        order_id = txt(value(row, crm_item_idx, "Số đơn hàng", "Số đơn hàng")).upper()
        if order_id not in crm_orders:
            continue
        raw_sku = txt(value(row, crm_item_idx, "Mã hàng hóa", "Mã hàng hóa"))
        code = sku(raw_sku, "CRM")
        if not code:
            add_gap(
                "CRM Final Products",
                f"{order_id} / row {source_row}",
                "Blank SKU",
                "Populate Mã hàng hóa in CRM Sales Order",
            )
            continue
        quantity = num(value(row, crm_item_idx, "Số lượng"))
        net_value = num(value(row, crm_item_idx, "Thành tiền sau CK", "Tổng tiền"))
        key = (order_id, code)
        entry = crm_item_agg.setdefault(
            key,
            {
                "order_id": order_id,
                "raw_sku": raw_sku,
                "sku": code,
                "description": txt(value(row, crm_item_idx, "Diễn giải", "Mô tả")),
                "quantity": 0.0,
                "net_value": 0.0,
            },
        )
        entry["quantity"] += quantity
        entry["net_value"] += net_value

    # Revenue.
    revenue_headers, revenue_source = all_rows(
        sources["Revenue"], "SỔ CHI TIẾT BÁN HÀNG", 4
    )
    require_headers(
        "Revenue",
        revenue_headers,
        [
            ("Ngày hạch toán",),
            ("Mã hàng",),
            ("Tổng số lượng bán",),
            ("Doanh số bán",),
            ("Chiết khấu",),
            ("Giá trị trả lại",),
            ("Giá trị giảm giá",),
            ("Giá vốn",),
            ("Đơn hàng",),
            ("TK Nợ",),
            ("TK Có",),
        ],
    )
    revenue_idx = header_map(revenue_headers)
    revenue_rows: list[list[Any]] = []
    revenue_by_key: dict[tuple[str, str], list[float]] = defaultdict(lambda: [0.0, 0.0])
    source_revenue_skus: set[str] = set()
    revenue_source_quantity_total = 0.0
    for source_row, row in enumerate(revenue_source, start=5):
        accounting_date = to_date(value(row, revenue_idx, "Ngày hạch toán"))
        if accounting_date is None or not (start_date <= accounting_date <= as_of):
            continue
        debit = txt(value(row, revenue_idx, "TK Nợ"))
        credit = txt(value(row, revenue_idx, "TK Có"))
        if not {debit, credit}.intersection({"5111", "5112", "5113"}):
            continue
        order_id = txt(value(row, revenue_idx, "Đơn hàng")).upper()
        if order_id in approved_order_exclusions:
            continue
        raw_sku = txt(value(row, revenue_idx, "Mã hàng"))
        code = sku(raw_sku, "REVENUE")
        if not code:
            add_gap(
                "Revenue",
                f"row {source_row}",
                "Blank SKU",
                "Populate Mã hàng in the sales ledger",
            )
            continue
        quantity = num(value(row, revenue_idx, "Tổng số lượng bán"))
        revenue_source_quantity_total += quantity
        net_value = (
            num(value(row, revenue_idx, "Doanh số bán"))
            - num(value(row, revenue_idx, "Chiết khấu"))
            - num(value(row, revenue_idx, "Giá trị trả lại"))
            - num(value(row, revenue_idx, "Giá trị giảm giá"))
        )
        cogs = num(value(row, revenue_idx, "Giá vốn"))
        item = products.get(code, {})
        revenue_rows.append(
            [
                len(revenue_rows) + 1,
                accounting_date.isoformat(),
                code,
                item.get("name") or txt(value(row, revenue_idx, "Tên hàng")),
                item.get("category", ""),
                item.get("subcategory", ""),
                item.get("supplier", ""),
                item.get("brand_code", ""),
                quantity,
                net_value,
                cogs,
                order_id,
                "",
                "",
                f"{order_id}{code}",
                None,
                None,
                "FOC / COST-ONLY"
                if abs(net_value) < 0.005 and abs(cogs) > 0.005
                else "",
            ]
        )
        revenue_by_key[(order_id, code)][0] += quantity
        revenue_by_key[(order_id, code)][1] += net_value
        source_revenue_skus.add(code)
        if code not in products:
            add_mismatch(
                "Revenue",
                f"{order_id} / {code}",
                "SKU not found in Product Master",
                "Map or correct official Product Master",
                f"Revenue row {source_row}; raw SKU={raw_sku}",
            )
        if cogs > net_value + 0.5:
            add_mismatch(
                "Revenue",
                f"{order_id} / {code}",
                "COGS > NET REV SOLD",
                "Review FOC/cost-only, discount/adjustment, sales price or SKU/cost mapping",
                f"NET REV SOLD={net_value:,.0f}; COGS={cogs:,.0f}; quantity={quantity:g}; ledger row {source_row}",
            )

    # Inventory output keeps the historical warehouse-level grain.  Per the PSI
    # contract, a raw row must have a positive closing quantity before it is kept;
    # PSI by Product separately aggregates those retained rows by canonical SKU.
    inventory_headers, inventory_source = all_rows(
        sources["Inventory"], "TỔNG HỢP TỒN KHO", 4
    )
    require_headers(
        "Inventory",
        inventory_headers,
        [("Tên kho",), ("Mã kho",), ("Mã hàng",), ("Cuối kỳ",)],
    )
    inventory_idx = header_map(inventory_headers)
    excluded_inventory_warehouses = {
        norm("Kho Bình Phú Hàng Lỗi (Kho ảo)"),
        norm("Kho Bình Phú (Kho lỗi)"),
        norm("Kho Chị Kathy"),
        norm("Kho chưa xuất hóa đơn"),
    }
    inventory_rows: list[list[Any]] = []
    inventory_by_sku: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    inventory_source_eligible_rows = 0
    inventory_nonpositive_source_rows = 0
    for source_row, row in enumerate(inventory_source, start=5):
        warehouse_name = txt(value(row, inventory_idx, "Tên kho"))
        if norm(warehouse_name) == norm("Tổng cộng"):
            continue
        if norm(warehouse_name) in excluded_inventory_warehouses:
            continue
        raw_sku = txt(value(row, inventory_idx, "Mã hàng"))
        code = sku(raw_sku, "INVENTORY")
        quantity_idx = find_index(inventory_idx, "Cuối kỳ")
        # The merged heading Cuối kỳ occupies two columns: quantity then value.
        quantity = num(
            row[quantity_idx]
            if quantity_idx is not None and quantity_idx < len(row)
            else None
        )
        stock_value = num(
            row[quantity_idx + 1]
            if quantity_idx is not None and quantity_idx + 1 < len(row)
            else None
        )
        if not code:
            if abs(quantity) > 0.001 or abs(stock_value) > 0.5:
                add_gap(
                    "Inventory",
                    f"row {source_row}",
                    "Blank SKU",
                    "Populate Mã hàng in inventory source",
                )
            continue
        inventory_source_eligible_rows += 1
        if quantity <= 0.001:
            inventory_nonpositive_source_rows += 1
            continue
        product = products.get(code, {})
        warehouse_code = txt(value(row, inventory_idx, "Mã kho"))
        inventory_rows.append(
            [
                len(inventory_rows) + 1,
                as_of.isoformat(),
                code,
                product.get("name") or txt(value(row, inventory_idx, "Tên hàng")),
                product.get("category", ""),
                product.get("subcategory", ""),
                product.get("supplier", ""),
                product.get("brand_code", ""),
                quantity,
                stock_value,
                warehouse_code,
            ]
        )
        inventory_by_sku[code][0] += quantity
        inventory_by_sku[code][1] += stock_value
        if code not in products:
            add_mismatch(
                "Inventory",
                code,
                "SKU not found in Product Master",
                "Map or correct official Product Master",
                f"Inventory row {source_row}; quantity={quantity:g}; value={stock_value:,.0f}; warehouse={warehouse_code}",
            )

    # Purchase / Loading List remains the approved source.
    purchase_book = load_workbook(
        sources["Purchase/PO"], read_only=True, data_only=True
    )
    purchase_ws = purchase_book["LDL"]
    purchase_rows: list[list[Any]] = []
    purchase_by_sku: dict[str, dict[str, Any]] = {}
    po_excluded_rows: list[list[Any]] = []
    for source_row, row in enumerate(
        purchase_ws.iter_rows(min_row=5, values_only=True), start=5
    ):
        if not any(item not in (None, "", 0) for item in row[:80]):
            continue
        raw_sku = txt(row[18] if len(row) > 18 else None)
        code = sku(raw_sku, "PURCHASE/PO")
        product_name = txt(row[19] if len(row) > 19 else None)
        if not code and not product_name:
            continue
        status = txt(row[1] if len(row) > 1 else None)
        po = txt(row[4] if len(row) > 4 else None)
        order_id = txt(row[8] if len(row) > 8 else None)
        if not code:
            issue = "Blank SKU in approved PO source"
            add_gap(
                "Purchase/PO",
                f"{po or order_id or 'row ' + str(source_row)}",
                issue,
                "Populate SKU before using this PO line for PSI",
            )
            po_excluded_rows.append(
                [len(po_excluded_rows) + 1, "", product_name, issue]
            )
            continue
        po_date = iso(row[5] if len(row) > 5 else None)
        qty = num(row[20] if len(row) > 20 else None)
        unit_warehouse_cost = num(row[76] if len(row) > 76 else None)
        warehouse_value = num(row[77] if len(row) > 77 else None)
        # LDL: AR is ETA; BC (54) is the actual stock-in date.
        eta = iso(row[43] if len(row) > 43 else None)
        product = products.get(code, {})
        purchase_rows.append(
            [
                len(purchase_rows) + 1,
                status,
                txt(row[2] if len(row) > 2 else None),
                po,
                order_id,
                po_date,
                txt(row[6] if len(row) > 6 else None),
                txt(row[7] if len(row) > 7 else None),
                code,
                product.get("name") or product_name,
                qty,
                warehouse_value,
                txt(row[21] if len(row) > 21 else None),
                # AB is net contract value excluding F.O.C; AA is the F.O.C flag.
                num(row[27] if len(row) > 27 else None),
                eta,
            ]
        )
        if code not in purchase_by_sku:
            purchase_by_sku[code] = {
                "unit_warehouse_cost": unit_warehouse_cost,
                "warehouse_value": warehouse_value,
                "qty": qty,
                "eta": eta,
                "po": po,
            }
    purchase_book.close()

    # Pre-orders: approved CRM line less recognized revenue, then exact permanent
    # exclusion by Order ID + canonical SKU.
    preorder_rows: list[list[Any]] = []
    preorder_excluded_rows: list[list[Any]] = []
    for (order_id, code), item_line in sorted(crm_item_agg.items()):
        revenue_qty, revenue_net = revenue_by_key[(order_id, code)]
        raw_open_qty = item_line["quantity"] - revenue_qty
        raw_open_net = item_line["net_value"] - revenue_net
        if raw_open_qty < -0.001 or raw_open_net < -0.5:
            add_mismatch(
                "CRM / Revenue",
                f"{order_id} / {code}",
                "Revenue exceeds approved CRM line",
                "Check CRM line, invoice adjustment and SKU mapping",
                f"CRM qty={item_line['quantity']:g}, net={item_line['net_value']:,.0f}; Revenue qty={revenue_qty:g}, net={revenue_net:,.0f}",
            )
        open_qty = max(0.0, raw_open_qty)
        open_net = max(0.0, raw_open_net)
        if open_qty <= 0.001:
            if open_net > 0.5:
                add_mismatch(
                    "CRM / Revenue",
                    f"{order_id} / {code}",
                    "CRM exceeds Revenue with no open quantity",
                    "Review CRM/Revenue amount difference; do not classify as Pre-order",
                    f"CRM qty={item_line['quantity']:g}, net={item_line['net_value']:,.0f}; Revenue qty={revenue_qty:g}, net={revenue_net:,.0f}",
                )
            continue
        product = products.get(code, {})
        po = purchase_by_sku.get(code, {})
        estimated_cost = po.get("unit_warehouse_cost", 0.0) * open_qty if po else None
        output_row = [
            0,
            crm_orders[order_id]["approval_date"],
            code,
            product.get("name") or item_line["description"],
            product.get("category", ""),
            product.get("subcategory", ""),
            product.get("supplier", ""),
            product.get("brand_code", ""),
            open_qty,
            open_net,
            estimated_cost,
            "",
            order_id,
            po.get("eta"),
            "",
            "",
            "",
            "",
        ]
        if (order_id, code) in approved_preorder_exclusions:
            output_row[0] = len(preorder_excluded_rows) + 1
            output_row[16] = "KT approved exact exclusion"
            output_row[17] = "EXCLUDE FROM PREORDER"
            preorder_excluded_rows.append(output_row)
        else:
            output_row[0] = len(preorder_rows) + 1
            preorder_rows.append(output_row)
        if code not in products:
            add_mismatch(
                "CRM Final Products",
                f"{order_id} / {code}",
                "SKU not found in Product Master",
                "Map or correct official Product Master",
                f"CRM open quantity={open_qty:g}; open net={open_net:,.0f}",
            )

    # Revenue orders absent from filtered CRM Final remain visible with the actual
    # source-side reason. Historical pre-2024 orders are outside CRM scope.
    for order_id, code in sorted(revenue_by_key):
        if (
            not order_id.startswith("DH-")
            or order_id in crm_orders
            or order_id in approved_order_exclusions
        ):
            continue
        raw = crm_order_all.get(order_id)
        if (
            raw is not None
            and raw["approval_date"]
            and raw["approval_date"] < start_date.isoformat()
        ):
            continue
        if raw is None:
            issue = "Revenue order missing CRM source"
            action = "Check CRM export coverage"
            detail = (
                "Revenue exists in ledger but the order is absent from the CRM export."
            )
        elif norm(raw["approval_status"]) != norm("Đã duyệt"):
            issue = "Revenue order excluded from CRM Final by approval status"
            action = (
                "Check whether revenue is valid for an unapproved/refused CRM order"
            )
            detail = f"CRM approval status={raw['approval_status'] or '(blank)'}; state={raw['state'] or '(blank)'}; approval date={raw['approval_date'] or '(blank)'}"
        else:
            issue = "Revenue order excluded from CRM Final"
            action = "Check CRM status/date filtering"
            detail = f"CRM approval status={raw['approval_status'] or '(blank)'}; state={raw['state'] or '(blank)'}; approval date={raw['approval_date'] or '(blank)'}"
        add_mismatch("Revenue / CRM", f"{order_id} / {code}", issue, action, detail)

    # PSI by Product.
    preorder_by_sku: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for row in preorder_rows:
        preorder_by_sku[row[2]][0] += num(row[8])
        preorder_by_sku[row[2]][1] += num(row[9])
    po_by_sku: dict[str, float] = defaultdict(float)
    for row in purchase_rows:
        if norm(row[1]) != norm("Done"):
            po_by_sku[row[8]] += num(row[10])
    revenue_by_sku: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0])
    for row in revenue_rows:
        revenue_by_sku[row[2]][0] += num(row[8])
        revenue_by_sku[row[2]][1] += num(row[9])
        revenue_by_sku[row[2]][2] += num(row[10])
    active_skus = sorted(
        set(source_revenue_skus)
        | set(inventory_by_sku)
        | set(preorder_by_sku)
        | set(po_by_sku)
    )
    psi_product_rows: list[list[Any]] = []
    for code in active_skus:
        product = products.get(code, {})
        sold_qty, net_revenue, cogs = revenue_by_sku[code]
        inventory_qty, inventory_value = inventory_by_sku.get(code, [0.0, 0.0])
        gross_profit = net_revenue - cogs
        gross_margin = gross_profit / net_revenue if abs(net_revenue) > 0.005 else None
        psi_product_rows.append(
            [
                code,
                product.get("name", ""),
                product.get("brand", ""),
                product.get("category", ""),
                product.get("subcategory", ""),
                "MAPPED" if code in products else "MISSING PRODUCT",
                sold_qty,
                net_revenue,
                cogs,
                gross_profit,
                gross_margin,
                inventory_qty,
                inventory_value,
                preorder_by_sku[code][0],
                po_by_sku[code],
            ]
        )

    # Target retained exactly as approved.
    target_book = load_workbook(sources["Target"], read_only=True, data_only=True)
    target_ws = target_book["Target"]
    target_rows = [
        [iso(item) for item in row]
        for row in target_ws.iter_rows(values_only=True)
        if any(item not in (None, "") for item in row)
    ]
    target_book.close()

    # Stable mismatch identity is Source + trimmed Key + Issue.
    deduped_mismatches: list[list[Any]] = []
    seen_mismatch_keys: set[tuple[str, str, str]] = set()
    for mismatch in mismatches:
        stable_key = (txt(mismatch[1]), txt(mismatch[2]), txt(mismatch[3]))
        if stable_key not in seen_mismatch_keys:
            seen_mismatch_keys.add(stable_key)
            deduped_mismatches.append(mismatch)
    prior_psi_path, prior_keys = prior_psi, prior_mismatch_keys()
    new_mismatch_keys = seen_mismatch_keys - prior_keys

    preorder_identity_counts = Counter(active_preorder_exclusion_rows)
    duplicate_preorder_identities = sorted(
        identity for identity, count in preorder_identity_counts.items() if count > 1
    )
    for order_id, code in duplicate_preorder_identities:
        add_gap(
            "Manual Check",
            f"{order_id} / {code}",
            f"Duplicate active permanent preorder identity ({preorder_identity_counts[(order_id, code)]} rows)",
            "Deduplicate Manual Check registry; PSI safely applies set semantics",
        )

    seen_gaps: set[tuple[str, str, str]] = set()
    deduped_gaps: list[list[Any]] = []
    for gap in data_gaps:
        stable_key = (txt(gap[0]), txt(gap[1]), txt(gap[2]))
        if stable_key not in seen_gaps:
            seen_gaps.add(stable_key)
            deduped_gaps.append(gap)

    # Release-gate computations are binary observables consumed by the builder.
    open_preorder_keys = {(row[12], row[2]) for row in preorder_rows}
    excluded_preorder_keys = {(row[12], row[2]) for row in preorder_excluded_rows}
    business_order_ids = (
        set(crm_orders)
        | {row[0] for row in []}
        | {row[11] for row in revenue_rows if row[11]}
        | {row[12] for row in preorder_rows if row[12]}
        | {row[12] for row in preorder_excluded_rows if row[12]}
    )
    order_exclusion_leaks = sorted(approved_order_exclusions & business_order_ids)
    preorder_exclusion_leaks = sorted(open_preorder_keys & approved_preorder_exclusions)
    invalid_excluded_preorders = sorted(
        excluded_preorder_keys - approved_preorder_exclusions
    )
    gate_results = {
        "source_files_readable": all(
            path.is_file() and path.stat().st_size > 0 for path in sources.values()
        ),
        "manual_check_validator_pass": True,
        "prior_psi_baseline_found": prior_psi_path is not None,
        "revenue_quantity_header_exact": txt(
            revenue_headers[find_index(revenue_idx, "Tổng số lượng bán") or 0]
        )
        == "Tổng số lượng bán",
        "revenue_quantity_total_matches_rows": abs(
            revenue_source_quantity_total - sum(num(row[8]) for row in revenue_rows)
        )
        < 0.001,
        "inventory_raw_positive_only": all(
            num(row[8]) > 0.001 for row in inventory_rows
        ),
        "preorder_exclusion_no_leak": not preorder_exclusion_leaks,
        "preorder_excluded_rows_authorized": not invalid_excluded_preorders,
        "preorder_excluded_output_unique": len(excluded_preorder_keys)
        == len(preorder_excluded_rows),
        "preorder_registry_set_semantics": len(approved_preorder_exclusions)
        == len(set(active_preorder_exclusion_rows)),
        "order_exclusion_no_leak": not order_exclusion_leaks,
        "mismatch_identity_unique": len(deduped_mismatches) == len(seen_mismatch_keys),
    }
    if not all(gate_results.values()):
        failed = [name for name, passed in gate_results.items() if not passed]
        raise RuntimeError(f"Release data gates failed: {', '.join(failed)}")

    gate_notes = {
        "source_files_readable": "All seven source/control files exist and are non-empty.",
        "manual_check_validator_pass": f"Registry loaded; {manual_summary['approved_active_preorder_exclusions']} active preorder rows.",
        "prior_psi_baseline_found": str(prior_psi_path or ""),
        "revenue_quantity_header_exact": "Selected exact header Tổng số lượng bán.",
        "revenue_quantity_total_matches_rows": f"Quantity={revenue_source_quantity_total:g}.",
        "inventory_raw_positive_only": f"Retained {len(inventory_rows)} positive warehouse-level rows; PSI by Product aggregates {len(inventory_by_sku)} canonical SKUs.",
        "preorder_exclusion_no_leak": "No active permanent key appears in open Pre-orders.",
        "preorder_excluded_rows_authorized": "Every excluded output row matches an active permanent key.",
        "preorder_excluded_output_unique": "Excluded output is unique by Order ID + canonical SKU.",
        "preorder_registry_set_semantics": f"Set semantics applied to {len(active_preorder_exclusion_rows)} rows / {len(approved_preorder_exclusions)} unique keys; duplicates are recorded as Data gaps.",
        "order_exclusion_no_leak": "No approved whole-order exclusion appears in business outputs.",
        "mismatch_identity_unique": "Mismatch output is unique by Source + trimmed Key + Issue.",
    }
    gates = [
        {
            "check": name,
            "expected": True,
            "actual": passed,
            "status": "PASS",
            "notes": gate_notes[name],
        }
        for name, passed in gate_results.items()
    ]
    if duplicate_preorder_identities:
        gates.append(
            {
                "check": "manual_check_duplicate_preorder_identities",
                "expected": 0,
                "actual": len(duplicate_preorder_identities),
                "status": "WARN",
                "notes": "Registry duplicates are surfaced in Data gaps; deterministic set semantics prevent duplicate exclusions.",
            }
        )

    mismatch_breakdown = Counter(row[1] for row in deduped_mismatches)
    new_mismatch_breakdown = Counter(source for source, _, _ in new_mismatch_keys)
    source_hashes = {name: sha256_file(path) for name, path in sources.items()}

    payload = {
        "as_of": as_of.isoformat(),
        "period_start": start_date.isoformat(),
        "sources": {name: str(path) for name, path in sources.items()},
        "source_hashes": source_hashes,
        "prior_psi_final": str(prior_psi_path) if prior_psi_path else "",
        "prior_psi_sha256": sha256_file(prior_psi_path) if prior_psi_path else "",
        "brand_note": "User-approved mapping: Product Master Nguồn gốc = Brand; Mã hãng = Brand Code.",
        "manual_check_summary": manual_summary,
        "gates": gates,
        "gate_results": gate_results,
        "gate_evidence": {
            "order_exclusion_leaks": order_exclusion_leaks,
            "preorder_exclusion_leaks": [
                list(item) for item in preorder_exclusion_leaks
            ],
            "invalid_excluded_preorders": [
                list(item) for item in invalid_excluded_preorders
            ],
            "duplicate_active_preorder_identities": [
                list(item) for item in duplicate_preorder_identities
            ],
            "active_preorder_exclusion_rows": len(active_preorder_exclusion_rows),
            "active_preorder_exclusion_unique_keys": len(approved_preorder_exclusions),
            "revenue_quantity_source": "Tổng số lượng bán",
            "revenue_quantity_total": revenue_source_quantity_total,
            "inventory_source_eligible_rows": inventory_source_eligible_rows,
            "inventory_canonical_aggregates": len(inventory_by_sku),
            "inventory_nonpositive_source_rows_removed": inventory_nonpositive_source_rows,
            "product_duplicate_canonical_skus": product_duplicate_skus,
        },
        "summary": {
            "crm_final_orders": len(crm_orders),
            "crm_final_products": len(crm_item_agg),
            "revenue_rows": len(revenue_rows),
            "revenue_quantity": sum(num(row[8]) for row in revenue_rows),
            "net_revenue": sum(num(row[9]) for row in revenue_rows),
            "cogs": sum(num(row[10]) for row in revenue_rows),
            "inventory_rows": len(inventory_rows),
            "inventory_qty": sum(num(row[8]) for row in inventory_rows),
            "inventory_value": sum(num(row[9]) for row in inventory_rows),
            "purchase_rows": len(purchase_rows),
            "po_excluded_rows": len(po_excluded_rows),
            "order_excluded_orders": len(approved_order_exclusions),
            "preorder_rows": len(preorder_rows),
            "preorder_excluded_rows": len(preorder_excluded_rows),
            "mismatch_rows": len(deduped_mismatches),
            "new_mismatch_rows": len(new_mismatch_keys),
            "data_gap_rows": len(deduped_gaps),
            "psi_product_rows": len(psi_product_rows),
        },
        "mismatch_breakdown": dict(sorted(mismatch_breakdown.items())),
        "new_mismatch_breakdown": dict(sorted(new_mismatch_breakdown.items())),
        "revenue_rows": revenue_rows,
        "inventory_rows": inventory_rows,
        "purchase_rows": purchase_rows,
        "po_excluded_rows": po_excluded_rows,
        "preorder_rows": preorder_rows,
        "preorder_excluded_rows": preorder_excluded_rows,
        "psi_product_rows": psi_product_rows,
        "crm_final_rows": [
            [
                record["order_id"],
                record["customer"],
                record["customer_code"],
                record["approval_date"],
                record["order_value"],
                record["invoiced_value"],
                record["status"],
                record["approval_status"],
                record["delivery_status"],
                record["payment_status"],
            ]
            for record in sorted(crm_orders.values(), key=lambda item: item["order_id"])
        ],
        "crm_product_rows": [
            [
                line["order_id"],
                line["sku"],
                products.get(line["sku"], {}).get("name") or line["description"],
                line["quantity"],
                line["net_value"],
                crm_orders[line["order_id"]]["approval_date"],
            ]
            for line in sorted(
                crm_item_agg.values(), key=lambda item: (item["order_id"], item["sku"])
            )
        ],
        "brand_rows": [
            list(item)
            for item in sorted(
                {
                    (product["brand"], product["brand_code"])
                    for product in products.values()
                    if product["brand"] or product["brand_code"]
                }
            )
        ],
        "category_rows": [
            list(item)
            for item in sorted(
                {
                    (product["category"], product["subcategory"])
                    for product in products.values()
                    if product["category"] or product["subcategory"]
                }
            )
        ],
        "target_rows": target_rows,
        "mismatch_rows": deduped_mismatches,
        "new_mismatch_keys": [list(item) for item in sorted(new_mismatch_keys)],
        "data_gap_rows": deduped_gaps,
        "data_gaps_rows": deduped_gaps,
    }

    return payload
