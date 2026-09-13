"""Validate a user-supplied offline Final before using it for NEW comparison."""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

SHEET_ORDER = (
    "PSI Summary",
    "Checks",
    "Sources",
    "Mismatch",
    "Data gaps",
    "PO excluded",
    "PSI by Product",
    "Brand",
    "Category",
    "Purchase",
    "Inventory",
    "Pre-orders",
    "Preorder excluded KT",
    "Revenue",
    "CRM Final",
    "Final CRM Products",
    "Target",
)
MISMATCH_HEADERS = (
    "Status",
    "Source",
    "Key",
    "Issue",
    "Resolution",
    "Evidence",
    "New since prior PSI",
)


def verify_baseline(content: bytes, filename: str, cutoff: date) -> date:
    """Require the established Final schema and a content date before cutoff.

    This accepts an offline Final explicitly selected by the operator. It does
    not manufacture a cloud approval receipt; published baselines must later be
    resolved through the authenticated release registry by immutable hash.
    """
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    if not re.search(r"\bfinal\b", name.replace("_", " "), re.IGNORECASE):
        raise ValueError("BASELINE_FINAL_REQUIRED")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if tuple(workbook.sheetnames) != SHEET_ORDER:
            raise ValueError("BASELINE_SCHEMA_INVALID")
        title = str(workbook["PSI Summary"]["A1"].value or "")
        match = re.fullmatch(r"PSI SUMMARY\s*[—–-]\s*(\d{2})\.(\d{2})\.(\d{4})", title)
        if not match and title != "PSI SUMMARY":
            raise ValueError("BASELINE_DATE_MISSING")
        title_date = None
        if match:
            day, month, year = map(int, match.groups())
            title_date = date(year, month, day)
        field_date = None
        if workbook["PSI Summary"]["A4"].value == "As of":
            raw = workbook["PSI Summary"]["B4"].value
            if isinstance(raw, datetime):
                field_date = raw.date()
            elif isinstance(raw, date):
                field_date = raw
            elif isinstance(raw, str):
                field_date = date.fromisoformat(raw)
            else:
                raise ValueError("BASELINE_DATE_MISSING")
        if title_date and field_date and title_date != field_date:
            raise ValueError("BASELINE_DATE_MISMATCH")
        prior_date = field_date or title_date
        if prior_date is None:
            raise ValueError("BASELINE_DATE_MISSING")
        if prior_date >= cutoff:
            raise ValueError("BASELINE_NOT_EARLIER")
        named = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", name)
        if named:
            day, month, year = map(int, named.groups())
            if date(year, month, day) != prior_date:
                raise ValueError("BASELINE_DATE_MISMATCH")
        headers = next(
            workbook["Mismatch"].iter_rows(
                min_row=3, max_row=3, max_col=7, values_only=True
            )
        )
        if headers != MISMATCH_HEADERS:
            raise ValueError("BASELINE_SCHEMA_INVALID")
        statuses = workbook["Checks"].iter_rows(
            min_row=4, min_col=4, max_col=4, values_only=True
        )
        values = [row[0] for row in statuses if row[0] is not None]
        if not values or any(value not in {"PASS", "WARN"} for value in values):
            raise ValueError("BASELINE_CHECKS_FAILED")
        return prior_date
    finally:
        workbook.close()
