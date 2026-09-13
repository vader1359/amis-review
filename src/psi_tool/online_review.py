"""Evidence-only review of a report and its explicitly selected prior Final.

Rules surface inconsistencies for accounting review; absence of a rule flag is
not approval. This module never changes report data or infers business causes.
"""

from __future__ import annotations

# Heterogeneous validated workbook cells and explicit API error codes.
# ruff: noqa: ANN401, EM101
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook

from ._online_workbook_schema import SCHEMA
from .online_baseline import verify_baseline

IDENTITY_PARTS = 2
QUANTITY_TOLERANCE = 0.001
MONEY_TOLERANCE = 0.5

# These describe the conditions actually raised by _online_prepare; suggested
# checks are possibilities to investigate, never asserted business causes.
ISSUE_GUIDANCE = {
    "COGS > NET REV SOLD": (
        (
            "Giá vốn cao hơn doanh thu thuần trên dòng bán hàng. "
            "Chênh lệch này chưa đủ để kết luận nhập sai hoặc bán lỗ ngoài dự kiến."
        ),
        (
            "Đối chiếu chứng từ để kiểm tra hàng tặng/chỉ ghi giá vốn, chiết khấu, "
            "điều chỉnh, giá bán và ánh xạ mã hàng/giá vốn."
        ),
    ),
    "Revenue exceeds approved CRM line": (
        (
            "Số lượng hoặc giá trị đã ghi nhận doanh thu vượt dòng CRM đã duyệt "
            "của cùng đơn và mã hàng. Cần đối chiếu để xác định bên nào cần điều chỉnh."
        ),
        (
            "So sánh số lượng và giá trị CRM với chứng từ bán hàng; kiểm tra "
            "điều chỉnh hóa đơn và ánh xạ mã hàng trước khi sửa dữ liệu."
        ),
    ),
    "CRM exceeds Revenue with no open quantity": (
        (
            "CRM còn giá trị lớn hơn doanh thu đã ghi nhận nhưng không còn số lượng "
            "chờ giao. Đây là chênh lệch tiền cần kiểm tra, "
            "chưa đủ căn cứ tạo preorder."
        ),
        (
            "Đối chiếu giá trị CRM và doanh thu, kiểm tra chiết khấu hoặc điều chỉnh; "
            "không đưa riêng phần chênh lệch tiền này vào preorder."
        ),
    ),
    "SKU not found in Product Master": (
        (
            "Mã hàng sau ánh xạ chưa có trong danh mục sản phẩm chính thức. "
            "Chưa xác định là thiếu danh mục hay cần sửa ánh xạ mã."
        ),
        (
            "Đối chiếu mã gốc với danh mục sản phẩm; bổ sung hoặc sửa mã/ánh xạ "
            "theo nguồn đã được duyệt."
        ),
    ),
    "Revenue order missing CRM source": (
        (
            "Đơn có doanh thu trong sổ bán hàng nhưng không xuất hiện trong file CRM "
            "đã cung cấp. Chưa thể kết luận đơn không tồn tại trong hệ thống CRM."
        ),
        (
            "Kiểm tra phạm vi và bộ lọc của file xuất CRM, sau đó tra cứu đơn "
            "trong hệ thống để xác minh dữ liệu còn thiếu."
        ),
    ),
    "Revenue order excluded from CRM Final by approval status": (
        (
            "Đơn có doanh thu nhưng bị loại khỏi CRM Final vì trạng thái phê duyệt "
            "không phải Đã duyệt. Trạng thái này chưa chứng minh doanh thu bị sai."
        ),
        (
            "Kiểm tra trạng thái phê duyệt thực tế và chứng từ để xác nhận việc "
            "ghi doanh thu cho đơn chưa được duyệt hoặc bị từ chối."
        ),
    ),
    "Revenue order excluded from CRM Final": (
        (
            "Đơn có doanh thu và có trong nguồn CRM nhưng không thuộc phạm vi "
            "CRM Final. Cần xác minh điều kiện lọc trạng thái hoặc ngày duyệt."
        ),
        (
            "Đối chiếu trạng thái, ngày duyệt và kỳ báo cáo với quy tắc lọc "
            "CRM Final trước khi điều chỉnh."
        ),
    ),
}


def _text(value: Any) -> str:
    return str(value or "").strip()


def _number(value: Any) -> float:
    return float(Decimal(str(value or 0)))


def _id(*parts: Any) -> str:
    return hashlib.sha256(
        json.dumps(parts, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()[:24]


def _value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return (
            value.date().isoformat()
            if isinstance(value, datetime)
            else value.isoformat()
        )
    return value


def _read_baseline(payload: dict[str, Any], content: bytes) -> dict[str, Any]:
    expected = payload.get("prior_psi_sha256")
    if not expected or hashlib.sha256(content).hexdigest() != expected:
        raise ValueError("REVIEW_BASELINE_HASH_MISMATCH")
    cutoff = date.fromisoformat(payload["as_of"])
    prior_date = verify_baseline(content, payload["prior_psi_final"], cutoff)
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    result: dict[str, Any] = {"as_of": prior_date.isoformat()}
    keys = {
        "crm_final_rows",
        "crm_product_rows",
        "revenue_rows",
        "preorder_rows",
        "preorder_excluded_rows",
    }
    try:
        for sheet in SCHEMA:
            if sheet["key"] not in keys:
                continue
            ws = book[sheet["name"]]
            width = len(sheet["headers"])
            headers = next(
                ws.iter_rows(min_row=3, max_row=3, max_col=width, values_only=True)
            )
            if list(headers) != sheet["headers"]:
                raise ValueError("REVIEW_BASELINE_SCHEMA_INVALID")
            result[sheet["key"]] = [
                [_value(value) for value in row]
                for row in ws.iter_rows(min_row=4, max_col=width, values_only=True)
                if any(value is not None for value in row)
            ]
    finally:
        book.close()
    return result


def _lines(payload: dict[str, Any]) -> dict[tuple[str, str], dict[str, Any]]:
    orders = {_text(row[0]): row for row in payload.get("crm_final_rows", [])}
    lines: dict[tuple[str, str], dict[str, Any]] = {}

    def line(order: Any, sku: Any) -> dict[str, Any]:
        key = (_text(order), _text(sku))
        if key not in lines:
            crm = orders.get(key[0])
            lines[key] = {
                "crm_quantity": 0.0,
                "crm_net": 0.0,
                "revenue_quantity": 0.0,
                "revenue_net": 0.0,
                "cogs": 0.0,
                "open_quantity": 0.0,
                "open_net": 0.0,
                "crm_present": False,
                "preorder_present": False,
                "excluded_from_preorder": False,
                "approval_date": _value(crm[3]) if crm else None,
                "order_value": _number(crm[4]) if crm else None,
                "invoiced_value": _number(crm[5]) if crm else None,
                "state": _text(crm[6]) if crm else None,
                "approval_status": _text(crm[7]) if crm else None,
                "delivery_status": _text(crm[8]) if crm else None,
                "payment_status": _text(crm[9]) if crm else None,
            }
        return lines[key]

    for row in payload.get("crm_product_rows", []):
        item = line(row[0], row[1])
        item["crm_present"] = True
        item["crm_quantity"] += _number(row[3])
        item["crm_net"] += _number(row[4])
    for row in payload.get("revenue_rows", []):
        item = line(row[11], row[2])
        item["revenue_quantity"] += _number(row[8])
        item["revenue_net"] += _number(row[9])
        item["cogs"] += _number(row[10])
    for row in payload.get("preorder_rows", []):
        item = line(row[12], row[2])
        item["preorder_present"] = True
        item["open_quantity"] += _number(row[8])
        item["open_net"] += _number(row[9])
    for row in payload.get("preorder_excluded_rows", []):
        line(row[12], row[2])["excluded_from_preorder"] = True
    return lines


def analyze_report(  # noqa: C901 - One deterministic review orchestration.
    payload: dict[str, Any],
    prior_payload: dict[str, Any] | None = None,
    prior_workbook: bytes | None = None,
) -> dict[str, Any]:
    """Return deterministic findings without asserting unproven root causes.

    ``prior_payload`` is an explicit trusted caller-provided baseline, never a
    latest-report lookup. Prefer ``prior_workbook`` to verify the exact Final
    referenced by the current report's lineage.
    """
    if prior_payload is not None and prior_workbook is not None:
        raise ValueError("REVIEW_BASELINE_AMBIGUOUS")
    if prior_workbook is not None:
        prior_payload = _read_baseline(payload, prior_workbook)
    if prior_payload is not None and (
        not prior_payload.get("as_of") or prior_payload["as_of"] >= payload["as_of"]
    ):
        raise ValueError("REVIEW_BASELINE_NOT_EARLIER")
    current = _lines(payload)
    previous = _lines(prior_payload) if prior_payload is not None else None
    new_keys = {
        tuple(_text(v) for v in row) for row in payload.get("new_mismatch_keys", [])
    }
    mismatches = []
    for row in payload.get("mismatch_rows", []):
        status, source, key, issue, action, evidence = map(_text, row[:6])
        parts = key.split(" / ", 1)
        identity = (
            tuple(parts)
            if len(parts) == IDENTITY_PARTS and tuple(parts) in current
            else None
        )
        explanation, suggested_action = ISSUE_GUIDANCE.get(
            issue,
            (
                (
                    "Báo cáo ghi nhận một vấn đề cần đối chiếu với nguồn. "
                    "Chưa có quy tắc giải thích riêng để kết luận nguyên nhân."
                ),
                "Đối chiếu bằng chứng và hướng xử lý gốc với người phụ trách nguồn.",
            ),
        )
        mismatches.append(
            {
                "id": _id("mismatch", source, key, issue),
                "status": status,
                "source": source,
                "key": key,
                "order_id": identity[0] if identity else None,
                "sku": identity[1] if identity else None,
                "issue": issue,
                "evidence": evidence,
                "explanation": explanation,
                "suggested_action": suggested_action,
                "source_action": action,
                "is_new": (source, key, issue) in new_keys,
            }
        )

    def reasons(key: tuple[str, str], item: dict[str, Any]) -> list[str]:
        found = [m["issue"] for m in mismatches if (m["order_id"], m["sku"]) == key]
        if not all(key):
            found.append(
                "Thiếu Order ID hoặc SKU; chưa đủ định danh để duyệt loại trừ."
            )
        if item["preorder_present"]:
            if not item["crm_present"]:
                found.append("Preorder không có dòng CRM đã duyệt tương ứng.")
            if item["open_quantity"] <= 0:
                found.append("Số lượng preorder không dương.")
            expected_qty = max(0, item["crm_quantity"] - item["revenue_quantity"])
            expected_net = max(0, item["crm_net"] - item["revenue_net"])
            if abs(item["open_quantity"] - expected_qty) > QUANTITY_TOLERANCE:
                found.append("Số lượng preorder khác CRM trừ doanh thu đã ghi nhận.")
            if abs(item["open_net"] - expected_net) > MONEY_TOLERANCE:
                found.append("Giá trị preorder khác CRM trừ doanh thu đã ghi nhận.")
            if item["open_net"] <= 0:
                found.append(
                    "Preorder còn số lượng nhưng giá trị không dương; cần kiểm tra."
                )
        return sorted(set(found))

    changes = []
    if previous is not None:
        for key in sorted(current.keys() | previous.keys()):
            before, after = previous.get(key), current.get(key)
            fields = sorted(
                k
                for k in (before or after or {})
                if (before or {}).get(k) != (after or {}).get(k)
            )
            if not fields:
                continue
            flags = reasons(key, after) if after is not None else []
            changes.append(
                {
                    "id": _id("change", *key),
                    "order_id": key[0],
                    "sku": key[1],
                    "kind": "added"
                    if before is None
                    else "removed"
                    if after is None
                    else "changed",
                    "before": before,
                    "after": after,
                    "changed_fields": fields,
                    "needs_attention": bool(flags),
                    "reasons": flags,
                }
            )
    checks = []
    for key, item in sorted(current.items()):
        if not item["preorder_present"]:
            continue
        flags = reasons(key, item)
        checks.append(
            {
                "id": _id("preorder", *key),
                "order_id": key[0],
                "sku": key[1],
                "is_new": not previous.get(key, {}).get("preorder_present", False)
                if previous is not None
                else None,
                "verdict": "needs_review" if flags else "no_rule_flags",
                "reasons": flags,
                "evidence": item,
            }
        )
    return {
        "schema_version": 1,
        "baseline": {
            "status": "available" if previous is not None else "missing",
            "source": "selected_final_workbook"
            if prior_workbook is not None
            else "explicit_payload"
            if previous is not None
            else None,
            "as_of": prior_payload.get("as_of") if prior_payload is not None else None,
            "sha256": payload.get("prior_psi_sha256")
            if prior_workbook is not None
            else None,
        },
        "summary": {
            "mismatch_count": len(mismatches),
            "change_count": len(changes),
            "new_preorder_count": sum(c["is_new"] is True for c in checks)
            if previous is not None
            else None,
            "preorder_attention_count": sum(
                c["verdict"] == "needs_review" for c in checks
            ),
        },
        "mismatches": mismatches,
        "changes": changes,
        "preorder_checks": checks,
    }
