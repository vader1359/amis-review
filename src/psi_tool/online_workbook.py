"""Portable PSI export and independent OOXML release validation.

Input is a validated business payload. Only the formulas authored here may be
written as formulas; all input strings (including Excel-like strings) stay text.
"""

# JSON payloads and third-party workbook APIs have dynamic value types.
# pyright: reportAny=false, reportExplicitAny=false, reportUnknownMemberType=false
# pyright: reportUnusedCallResult=false, reportImplicitStringConcatenation=false
# pyright: reportUnnecessaryIsInstance=false, reportMissingTypeStubs=false
# Column indices and format literals mirror the approved 17-sheet workbook.
# ruff: noqa: PLR2004, C901, PLR0912, PLR0915
from __future__ import annotations

import io
import math
import re
import zipfile
from collections.abc import Iterable  # noqa: TC003 - runtime API annotation
from datetime import UTC, date, datetime
from typing import Any, cast
from xml.etree import ElementTree as ET

import openpyxl
import xlsxwriter
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from xlsxwriter.format import Format  # noqa: TC002 - runtime API annotation

from psi_tool._online_workbook_schema import SCHEMA

MAX_ROWS = 250_000
MAX_CELLS = 2_000_000
MAX_ZIP_BYTES = 128 * 1024 * 1024
MAX_XML_BYTES = 512 * 1024 * 1024
NUM_FMT = "#,##0.0;[Red](#,##0.0);-"
MONEY_FMT = "#,##0;[Red](#,##0);-"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _filename(value: object) -> str:
    return str(value or "").replace("\\", "/").rsplit("/", 1)[-1]


def _key(row: Iterable[object]) -> tuple[str, ...]:
    return tuple(str(v or "").strip() for v in row)


def _note(name: str, payload: dict[str, Any]) -> str:
    if name == "Mismatch":
        prior = payload.get("prior_as_of", "")
        if not prior:
            match = re.search(
                r"(\d{2})[.](\d{2})[.](\d{4})",
                _filename(payload.get("prior_psi_final")),
            )
            prior = (
                match.group(0) if match else _filename(payload.get("prior_psi_final"))
            )
        elif re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(prior)):
            prior = date.fromisoformat(prior).strftime("%d.%m.%Y")
        return (
            f"Dòng NEW là mismatch mới so với PSI Final {prior}; "
            "trạng thái OPEN được giữ để xử lý, không tự suppress."
        )
    if name == "Inventory":
        return (
            "Giữ format theo kho như PSI lịch sử: chuẩn hoá SKU, "
            "loại kho không đủ điều kiện và chỉ giữ từng dòng có tồn cuối kỳ dương. "
            "PSI by Product aggregate riêng theo SKU."
        )
    if name == "Checks":
        return (
            "Các gate release được tính độc lập trong bước chuẩn bị dữ liệu. "
            "Workbook sẽ không xuất nếu một gate FAIL."
        )
    if name == "Brand":
        return payload.get(
            "brand_note",
            "Brand được lấy từ Product Master; SKU/Brand không xác định được "
            "giữ tại Data gaps để xử lý, không suy diễn mapping.",
        )
    if name == "Target":
        return "Đã giới hạn phạm vi dữ liệu Target ở cột A:I."
    return (
        "Nguồn và quy tắc theo PSI process cập nhật; kỳ dữ liệu "
        f"{payload.get('period_start', '2024-01-01')} đến {payload['as_of']}."
    )


def _definitions(payload: dict[str, Any]) -> list[dict[str, Any]]:
    gates = cast("list[dict[str, Any]]", payload.get("gates"))
    if not isinstance(gates, list) or not gates:
        msg = "Release gates are required"
        raise ValueError(msg)
    if any(
        not isinstance(g, dict)
        or str(g.get("status", "")).upper() not in {"PASS", "WARN"}
        for g in gates
    ):
        msg = "Release gates contain FAIL or invalid status"
        raise ValueError(msg)
    _ = date.fromisoformat(payload["as_of"])
    sources = payload.get("sources", {})
    hashes = payload.get("source_hashes", {})
    if not sources or any(
        not re.fullmatch(r"[0-9a-fA-F]{64}", str(hashes.get(k, ""))) for k in sources
    ):
        msg = "Every source requires a SHA-256 hash"
        raise ValueError(msg)
    if payload.get("prior_psi_final") and not re.fullmatch(
        r"[0-9a-fA-F]{64}", str(payload.get("prior_psi_sha256", ""))
    ):
        msg = "Prior PSI requires a SHA-256 hash"
        raise ValueError(msg)
    new = {_key(k) for k in payload.get("new_mismatch_keys", [])}
    mismatches = [
        [*list(r), "NEW" if _key(r[1:4]) in new else ""]
        for r in payload.get("mismatch_rows", [])
    ]
    target = [list(r[:9]) for r in payload.get("target_rows", [])]
    special: dict[str, list[list[Any]]] = {
        "PSI Summary": [
            [
                "As of",
                payload["as_of"],
                "Data period",
                f"{payload.get('period_start', '2024-01-01')} → {payload['as_of']}",
            ],
            ["Net revenue", None, "COGS", None],
            ["Revenue quantity", None, "Revenue lines", None],
            ["Inventory quantity", None, "Inventory value", None],
            ["Open Pre-order lines", None, "PO lines", None],
            ["CRM Final orders", None, "Mismatch rows", None],
            [
                "New mismatch rows",
                None,
                "Prior PSI Final",
                _filename(payload.get("prior_psi_final")),
            ],
        ],
        "Checks": [
            [
                g.get("check", g.get("name", "")),
                g.get("expected", ""),
                g.get("actual", ""),
                g["status"],
                g.get("notes", g.get("note", "")),
            ]
            for g in gates
        ],
        "Sources": [
            [k, _filename(v), "Official input / approved control", hashes[k]]
            for k, v in sources.items()
        ]
        + [
            [
                "Prior PSI Final",
                _filename(payload.get("prior_psi_final")),
                "Baseline for NEW mismatch comparison",
                payload.get("prior_psi_sha256", ""),
            ]
        ],
        "Mismatch": mismatches,
        "Data gaps": payload.get("data_gap_rows")
        or [
            [r[1], r[2], r[3], r[4]]
            for r in mismatches
            if re.search(
                r"SKU not found|missing from CRM|blank SKU|missing SKU",
                str(r[3]),
                re.IGNORECASE,
            )
        ],
        "Target": target[1:],
    }
    definitions: list[dict[str, Any]] = []
    cells = 0
    for spec in SCHEMA:
        name = spec["name"]
        headers = (
            (target[0] if target else ["Target"])
            if name == "Target"
            else spec["headers"]
        )
        rows: list[list[Any]] = (
            special[name] if name in special else payload.get(spec["key"], [])
        )
        if not isinstance(rows, list) or len(rows) > MAX_ROWS:
            msg = f"{name}: invalid row collection or row limit exceeded"
            raise ValueError(msg)
        if any(
            not isinstance(row, (list, tuple)) or len(row) != len(headers)
            for row in rows
        ):
            msg = f"{name}: row width differs from schema"
            raise ValueError(msg)
        cells += len(rows) * len(headers)
        if cells > MAX_CELLS:
            msg = "Workbook cell limit exceeded"
            raise ValueError(msg)
        for row in [headers, *rows]:
            for value in row:
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    msg = f"{name}: unsupported cell value"
                    raise ValueError(msg)
                if isinstance(value, float) and not math.isfinite(value):
                    msg = f"{name}: non-finite number"
                    raise ValueError(msg)
                if isinstance(value, str) and (
                    len(value) > 32767
                    or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)
                ):
                    msg = f"{name}: invalid Excel text"
                    raise ValueError(msg)
        if name == "PSI by Product":
            for row in rows:
                for value in (row[7], row[8]):
                    if value not in (None, "") and not isinstance(value, (int, float)):
                        msg = "PSI formula inputs must be numeric or blank"
                        raise ValueError(msg)
                revenue, cost = row[7] or 0, row[8] or 0
                if not math.isfinite(revenue - cost) or (
                    revenue and not math.isfinite((revenue - cost) / revenue)
                ):
                    msg = "PSI formula result must be finite"
                    raise ValueError(msg)
        title = (
            f"PSI SUMMARY — {date.fromisoformat(payload['as_of']).strftime('%d.%m.%Y')}"
            if name == "PSI Summary"
            else spec["title"]
        )
        definitions.append(
            dict(
                spec, headers=headers, rows=rows, title=title, note=_note(name, payload)
            )
        )
    return definitions


def _summary_formulas(
    definitions: list[dict[str, Any]],
) -> dict[str, tuple[str, int | float]]:
    rows = {d["name"]: d["rows"] for d in definitions}
    result: dict[str, tuple[str, int | float]] = {}
    for address, name, col, operation in [
        ("B5", "Revenue", "J", "SUM"),
        ("D5", "Revenue", "K", "SUM"),
        ("B6", "Revenue", "I", "SUM"),
        ("D6", "Revenue", "A", "COUNTA"),
        ("B7", "Inventory", "I", "SUM"),
        ("D7", "Inventory", "J", "SUM"),
        ("B8", "Pre-orders", "A", "COUNTA"),
        ("D8", "Purchase", "A", "COUNTA"),
        ("B9", "CRM Final", "A", "COUNTA"),
        ("D9", "Mismatch", "A", "COUNTA"),
        ("B10", "Mismatch", "G", "COUNTIF"),
    ]:
        end = max(4, len(rows[name]) + 3)
        suffix = ',"NEW"' if operation == "COUNTIF" else ""
        formula = f"={operation}('{name}'!{col}4:{col}{end}{suffix})"
        values = [r[ord(col) - 65] for r in rows[name]]
        cached = (
            sum(
                v
                for v in values
                if isinstance(v, (int, float)) and not isinstance(v, bool)
            )
            if operation == "SUM"
            else sum(v not in (None, "") for v in values)
            if operation == "COUNTA"
            else values.count("NEW")
        )
        result[address] = formula, cached
    return result


def build_workbook(payload: dict[str, Any]) -> bytes:
    """Render validated PSI data into a portable, deterministic 17-sheet XLSX."""
    definitions = _definitions(payload)
    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(
        output,
        {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False},
    )
    workbook.set_properties(
        {"title": "PSI Final", "created": datetime(2000, 1, 1, tzinfo=UTC)}
    )

    def add_format(properties: dict[str, Any]) -> Format:
        return workbook.add_format(dict(properties, font_name="Carlito"))

    title_fmt = add_format(
        {
            "bg_color": "#17365D",
            "font_color": "#FFFFFF",
            "bold": True,
            "font_size": 14,
            "valign": "vcenter",
        }
    )
    note_fmt = add_format(
        {
            "bg_color": "#D9EAF7",
            "font_color": "#1F2937",
            "italic": True,
            "text_wrap": True,
            "valign": "vcenter",
        }
    )
    head_fmt = add_format(
        {
            "bg_color": "#1F4E78",
            "font_color": "#FFFFFF",
            "bold": True,
            "text_wrap": True,
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "border_color": "#17365D",
        }
    )
    base = {
        "font_size": 9,
        "font_color": "#1F2937",
        "valign": "vcenter",
        "bottom": 1,
        "bottom_color": "#D9E2F3",
    }
    formats = {
        k: add_format(dict(base, **({"num_format": v} if v else {})))
        for k, v in [
            ("text", None),
            ("quantity", NUM_FMT),
            ("money", MONEY_FMT),
            ("percent", "0.0%"),
        ]
    }
    red = add_format({"bg_color": "#FCE4D6", "font_color": "#9C0006", "bold": True})
    amber = add_format({"bg_color": "#FFF2CC", "font_color": "#9C6500", "bold": True})
    green = add_format({"bg_color": "#E2F0D9", "font_color": "#006100", "bold": True})
    summary_formulas = _summary_formulas(definitions)
    label_fmt = add_format(dict(base, font_size=11, bg_color="#D9EAF7", bold=True))
    summary_formats = {
        k: add_format({"font_size": 11, "valign": "vcenter", "num_format": v})
        for k, v in [("text", "General"), ("money", MONEY_FMT), ("quantity", NUM_FMT)]
    }
    for definition in definitions:
        name, headers, rows = (
            definition["name"],
            definition["headers"],
            definition["rows"],
        )
        sheet = workbook.add_worksheet(name)
        sheet.hide_gridlines(2)
        sheet.freeze_panes(3, 0)
        width = len(headers)
        for col in range(width):
            # XlsxWriter adds 5 pixels of padding. Reverse that conversion so
            # stored OOXML widths equal the offline renderer's exact widths.
            sheet.set_column(col, col, definition["widths"].get(col, 18) - 5 / 7)
        for row, text, fmt, height in [
            (0, definition["title"], title_fmt, 28),
            (1, definition["note"], note_fmt, 24),
        ]:
            if width > 1:
                sheet.merge_range(row, 0, row, width - 1, text, fmt)
            else:
                sheet.write_string(row, 0, text, fmt)
            sheet.set_row(row, height)
        sheet.write_row(2, 0, headers, head_fmt)
        sheet.set_row(2, 34)
        for row_index, values in enumerate(rows, 3):
            for col, value in enumerate(values):
                category = next(
                    (
                        k
                        for k in ("quantity", "money", "percent")
                        if col in definition[k]
                    ),
                    "text",
                )
                fmt = formats[category]
                if name == "PSI Summary":
                    fmt = (
                        label_fmt
                        if col in (0, 2)
                        else summary_formats[
                            "text"
                            if row_index == 3 or (row_index == 9 and col == 3)
                            else "quantity"
                            if row_index >= 5
                            else "money"
                        ]
                    )
                if name == "PSI by Product" and col in (9, 10):
                    revenue, cost = values[7], values[8]
                    rev = revenue if isinstance(revenue, (int, float)) else 0
                    cogs = cost if isinstance(cost, (int, float)) else 0
                    profit = rev - cogs
                    formula = (
                        f"=H{row_index + 1}-I{row_index + 1}"
                        if col == 9
                        else f"=IFERROR(J{row_index + 1}/H{row_index + 1},0)"
                    )
                    sheet.write_formula(
                        row_index,
                        col,
                        formula,
                        fmt,
                        cast("Any", profit if col == 9 else profit / rev if rev else 0),
                    )
                elif (
                    name == "PSI Summary"
                    and f"{get_column_letter(col + 1)}{row_index + 1}"
                    in summary_formulas
                ):
                    formula, cached = summary_formulas[
                        f"{get_column_letter(col + 1)}{row_index + 1}"
                    ]
                    sheet.write_formula(
                        row_index, col, formula, fmt, cast("Any", cached)
                    )
                else:
                    sheet.write(row_index, col, value if value is not None else "", fmt)
        end = len(rows) + 2
        if rows and name == "Revenue":
            sheet.conditional_format(
                3, 9, end, 9, {"type": "formula", "criteria": "=$K4>$J4", "format": red}
            )
        if rows and name == "Mismatch":
            sheet.conditional_format(
                3,
                0,
                end,
                0,
                {
                    "type": "text",
                    "criteria": "containing",
                    "value": "OPEN",
                    "format": amber,
                },
            )
            sheet.conditional_format(
                3,
                0,
                end,
                6,
                {"type": "formula", "criteria": '=$G4="NEW"', "format": red},
            )
        if rows and name == "Checks":
            for status, fmt in [("PASS", green), ("WARN", amber), ("FAIL", red)]:
                sheet.conditional_format(
                    3,
                    3,
                    end,
                    3,
                    {
                        "type": "text",
                        "criteria": "containing",
                        "value": status,
                        "format": fmt,
                    },
                )
    workbook.close()
    content = output.getvalue()
    if len(content) > MAX_ZIP_BYTES:
        msg = "Workbook byte limit exceeded"
        raise ValueError(msg)
    return content


def validate_workbook(content: bytes, payload: dict[str, Any]) -> dict[str, Any]:
    """Independently inspect exported values, formulas, layout and ZIP safety."""
    failures: list[str] = []
    failure_count = 0

    def fail(message: str) -> None:
        nonlocal failure_count
        failure_count += 1
        if len(failures) < 100:
            failures.append(message)

    try:
        definitions = _definitions(payload)
        if len(content) > MAX_ZIP_BYTES:
            msg = "Workbook byte limit exceeded"
            raise ValueError(msg)  # noqa: TRY301 - convert all invalid archives to FAIL
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_XML_BYTES:
                msg = "Workbook expanded byte limit exceeded"
                raise ValueError(msg)  # noqa: TRY301 - guarded package boundary
            if len(archive.namelist()) != len(set(archive.namelist())):
                fail("Duplicate ZIP entries")
            for entry in archive.namelist():
                if (
                    "externalLink" in entry
                    or "vbaProject" in entry
                    or entry.startswith("/")
                    or ".." in entry.split("/")
                ):
                    fail(f"Forbidden package entry: {entry}")
                if entry.endswith((".xml", ".rels")) and re.search(
                    rb"<!DOCTYPE|<!ENTITY", archive.read(entry), re.IGNORECASE
                ):
                    msg = "XML entity declarations are forbidden"
                    raise ValueError(msg)  # noqa: TRY301 - reject before any XML parser
                if entry.endswith(".rels"):
                    rels = ET.fromstring(archive.read(entry))  # noqa: S314 - declarations rejected above
                    if any(r.get("TargetMode") == "External" for r in rels):
                        fail("External relationship present")
            styles = ET.fromstring(archive.read("xl/styles.xml"))  # noqa: S314 - declarations rejected above
            differentials = styles.findall("s:dxfs/s:dxf", NS)
            sheets_xml = [
                ET.fromstring(archive.read(f"xl/worksheets/sheet{i}.xml"))  # noqa: S314 - declarations rejected above
                for i in range(1, 18)
            ]
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        if workbook.sheetnames != [d["name"] for d in definitions]:
            fail("Sheet order or sheet names differ from the 17-sheet schema")
        expected_summary = _summary_formulas(definitions)
        formula_count = 0
        for spec, xml in zip(definitions, sheets_xml, strict=True):
            name, rows, headers = spec["name"], spec["rows"], spec["headers"]
            if name not in workbook:
                continue
            sheet = workbook[name]
            width = len(headers)
            expected_merges: set[str] = (
                {f"A1:{get_column_letter(width)}1", f"A2:{get_column_letter(width)}2"}
                if width > 1
                else set()
            )
            if {
                m.get("ref") for m in xml.findall("s:mergeCells/s:mergeCell", NS)
            } != expected_merges:
                fail(f"{name}: title/note merges differ")
            pane = xml.find("s:sheetViews/s:sheetView/s:pane", NS)
            if (
                pane is None
                or pane.get("topLeftCell") != "A4"
                or pane.get("ySplit") != "3"
            ):
                fail(f"{name}: freeze pane must be A4")
            view = xml.find("s:sheetViews/s:sheetView", NS)
            if view is None or view.get("showGridLines") != "0":
                fail(f"{name}: gridlines must be hidden")
            layout_rows = xml.findall("s:sheetData/s:row", NS)
            for row_number, height in [(1, 28), (2, 24), (3, 34)]:
                row_element = next(
                    (r for r in layout_rows if r.get("r") == str(row_number)), None
                )
                if row_element is None or float(row_element.get("ht", "0")) != height:
                    fail(f"{name}: row {row_number} height differs")
            columns = xml.findall("s:cols/s:col", NS)
            for col in range(1, width + 1):
                dimension = next(
                    (
                        c
                        for c in columns
                        if int(c.attrib["min"]) <= col <= int(c.attrib["max"])
                    ),
                    None,
                )
                if (
                    dimension is None
                    or abs(
                        float(dimension.attrib["width"])
                        - spec["widths"].get(col - 1, 18)
                    )
                    > 0.01
                ):
                    fail(f"{name}: column {col} width differs")
            if sheet.max_row != max(3, len(rows) + 3) or sheet.max_column != width:
                fail(f"{name}: dimensions differ from payload")
            cached: dict[str, str | None] = {}
            for cell in xml.findall("s:sheetData/s:row/s:c", NS):
                address = cell.attrib["r"]
                cell_row, cell_col = coordinate_to_tuple(address)
                if cell_row > len(rows) + 3 or cell_col > width:
                    fail(f"{name}!{address}: cell outside declared schema")
                if cell.get("t") == "e":
                    fail(f"{name}!{address}: cached Excel error")
                formula = cell.find("s:f", NS)
                if formula is not None:
                    allowed = (
                        name == "PSI Summary" and address in expected_summary
                    ) or (
                        name == "PSI by Product"
                        and 4 <= cell_row <= len(rows) + 3
                        and cell_col in (10, 11)
                    )
                    if not allowed:
                        fail(f"{name}!{address}: unauthorized formula in XML")
                    value = cell.find("s:v", NS)
                    cached[address] = value.text if value is not None else None
            for row_index, cells in enumerate(sheet.iter_rows(), 1):
                for col, cell in enumerate(cells):
                    expected = (
                        spec["title"]
                        if row_index == 1 and col == 0
                        else spec["note"]
                        if row_index == 2 and col == 0
                        else headers[col]
                        if row_index == 3
                        else rows[row_index - 4][col]
                        if row_index >= 4 and row_index - 4 < len(rows)
                        else None
                    )
                    if cell.value is not None:
                        font_size = (
                            14
                            if row_index == 1
                            else 11
                            if row_index <= 3 or name == "PSI Summary"
                            else 9
                        )
                        if cell.font.name != "Carlito" or cell.font.sz != font_size:
                            fail(f"{name}!{cell.coordinate}: font differs")
                        if row_index <= 3:
                            color = {1: "FF17365D", 2: "FFD9EAF7", 3: "FF1F4E78"}[
                                row_index
                            ]
                            if cell.fill.fgColor.rgb != color:
                                fail(f"{name}!{cell.coordinate}: header fill differs")
                        number_format = "General"
                        if row_index >= 4:
                            if (
                                name == "PSI Summary"
                                and col in (1, 3)
                                and row_index >= 5
                                and not (row_index == 10 and col == 3)
                            ):
                                number_format = MONEY_FMT if row_index == 5 else NUM_FMT
                            elif name != "PSI Summary":
                                number_format = (
                                    NUM_FMT
                                    if col in spec["quantity"]
                                    else MONEY_FMT
                                    if col in spec["money"]
                                    else "0.0%"
                                    if col in spec["percent"]
                                    else "General"
                                )
                        if cell.number_format != number_format:
                            fail(f"{name}!{cell.coordinate}: number format differs")
                    authored = None
                    expected_cache: int | float = 0
                    if name == "PSI Summary" and cell.coordinate in expected_summary:
                        authored, expected_cache = expected_summary[cell.coordinate]
                    elif name == "PSI by Product" and row_index >= 4 and col in (9, 10):
                        authored = (
                            f"=H{row_index}-I{row_index}"
                            if col == 9
                            else f"=IFERROR(J{row_index}/H{row_index},0)"
                        )
                        source = rows[row_index - 4]
                        rev = source[7] if isinstance(source[7], (int, float)) else 0
                        cost = source[8] if isinstance(source[8], (int, float)) else 0
                        expected_cache = (
                            rev - cost if col == 9 else (rev - cost) / rev if rev else 0
                        )
                    if authored:
                        formula_count += 1
                        if cell.data_type != "f" or cell.value != authored:
                            fail(f"{name}!{cell.coordinate}: authored formula differs")
                        try:
                            if not math.isclose(
                                float(cached.get(cell.coordinate) or "nan"),
                                float(expected_cache),
                                rel_tol=1e-12,
                                abs_tol=1e-7,
                            ):
                                fail(f"{name}!{cell.coordinate}: formula cache differs")
                        except (ValueError, TypeError):
                            fail(
                                f"{name}!{cell.coordinate}: "
                                "formula cache missing or invalid"
                            )
                    elif cell.data_type == "f":
                        fail(f"{name}!{cell.coordinate}: unauthorized formula")
                    elif (cell.value if cell.value is not None else "") != (
                        expected if expected is not None else ""
                    ):
                        # OOXML stores numbers with 15 significant digit precision.
                        if not (
                            isinstance(expected, (int, float))
                            and isinstance(cell.value, (int, float))
                            and math.isclose(
                                cell.value, expected, rel_tol=1e-14, abs_tol=1e-8
                            )
                        ):
                            fail(
                                f"{name}!{cell.coordinate}: value differs from payload"
                            )
            if name == "Mismatch" and rows:
                rules = xml.findall("s:conditionalFormatting", NS)
                if not any(
                    r.get("sqref") == f"A4:G{len(rows) + 3}"
                    and any(
                        f.text in ('$G4="NEW"', '=$G4="NEW"')
                        for f in r.findall("s:cfRule/s:formula", NS)
                    )
                    for r in rules
                ):
                    fail("Mismatch: missing full-row NEW conditional formatting")
                for rule_group in rules:
                    if rule_group.get("sqref") != f"A4:G{len(rows) + 3}":
                        continue
                    for rule in rule_group.findall("s:cfRule", NS):
                        formulas = rule.findall("s:formula", NS)
                        if not any(
                            f.text in ('$G4="NEW"', '=$G4="NEW"') for f in formulas
                        ):
                            continue
                        differential = differentials[int(rule.attrib["dxfId"])]
                        color = differential.find("s:fill/s:patternFill/s:bgColor", NS)
                        if color is None:
                            color = differential.find(
                                "s:fill/s:patternFill/s:fgColor", NS
                            )
                        bold = differential.find("s:font/s:b", NS)
                        if (
                            color is None
                            or color.get("rgb") != "FFFCE4D6"
                            or bold is None
                            or bold.get("val", "1") == "0"
                        ):
                            fail(
                                "Mismatch: NEW conditional formatting "
                                "must be red and bold"
                            )
        workbook.close()
    except Exception as exc:  # noqa: BLE001 - public validator returns FAIL for malformed workbooks
        fail(f"Validation could not complete: {type(exc).__name__}: {exc}")
        formula_count = 0
    return {
        "status": "FAIL" if failure_count else "PASS",
        "failure_count": failure_count,
        "failures": failures,
        "sheet_count": 17 if not failure_count else None,
        "formula_count": formula_count,
    }
