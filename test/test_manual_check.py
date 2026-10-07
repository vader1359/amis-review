from __future__ import annotations

from datetime import date
from decimal import Decimal
from hashlib import sha256
import os
from pathlib import Path

from openpyxl import load_workbook
import pytest

from psi_engine.manual_check import (
    ExceptionRule,
    ManualCheckRegistry,
    OrderExclusion,
    PreorderExclusion,
    SkuMapping,
    exclude_orders,
    exclude_preorders,
    load_manual_check,
    make_fingerprint,
    make_order_exclusion_fingerprint,
    make_preorder_fingerprint,
    partition_mismatches,
)


AS_OF = date(2026, 8, 7)


def exception_rule(*, status: str = "APPROVED", action: str = "IGNORE MISMATCH", effective_to: date | None = None) -> ExceptionRule:
    issue_type = "CRM MISSING REVENUE"
    return ExceptionRule(
        exception_id="EXC-TEST",
        fingerprint=make_fingerprint(action=action, scope="ORDER", order_id="DH-0001", issue_type=issue_type),
        status=status,
        action=action,
        scope="ORDER",
        order_id="DH-0001",
        raw_sku="",
        canonical_sku="",
        quantity=None,
        value=None,
        issue_type=issue_type,
        reason="Approved exact order exception",
        evidence="unit test",
        effective_from=date(2024, 1, 1),
        effective_to=effective_to,
        approved_by="KT",
        approved_date=date(2026, 8, 1),
        notes="",
    )


def preorder_rule() -> PreorderExclusion:
    return PreorderExclusion(
        exclusion_id="PREX-TEST",
        fingerprint=make_preorder_fingerprint(
            action="EXCLUDE FROM PREORDER",
            order_id="DH-0002",
            canonical_sku="SKU.30",
        ),
        status="APPROVED",
        action="EXCLUDE FROM PREORDER",
        order_id="DH-0002",
        raw_sku="SKU.3",
        canonical_sku="SKU.30",
        quantity=Decimal("2"),
        net_value=Decimal("100"),
        source_row=4,
        source_no="1",
        kt_note="KT approved",
        treatment="Clear lại CRM, loại khỏi Pre",
        effective_from=date(2024, 1, 1),
        effective_to=None,
        approved_by="KT",
        approved_date=date(2026, 8, 1),
        evidence="unit test",
        notes="",
        disposition="PERMANENT / DONE",
    )


def order_rule() -> OrderExclusion:
    action = "EXCLUDE ORDER FROM PSI"
    scope = "ALL PSI BUSINESS SHEETS"
    return OrderExclusion(
        exclusion_id="OREX-TEST",
        fingerprint=make_order_exclusion_fingerprint(action=action, scope=scope, order_id="DH-CANCELLED"),
        status="APPROVED",
        action=action,
        scope=scope,
        order_id="DH-CANCELLED",
        reason="Cancelled order",
        treatment="Remove from PSI business sheets",
        effective_from=date(2024, 1, 1),
        effective_to=None,
        approved_by="IAN",
        approved_date=date(2026, 8, 11),
        evidence="unit test",
        notes="",
        disposition="PERMANENT / DONE",
    )


def suffix_mapping() -> SkuMapping:
    return SkuMapping(
        mapping_id="MAP-TEST",
        fingerprint=make_fingerprint(
            action="MAP SKU",
            scope="SKU",
            raw_sku=".3",
            canonical_sku=".30",
            issue_type="SUFFIX:ALL OFFICIAL SOURCES",
        ),
        status="APPROVED",
        action="MAP SKU",
        match_type="SUFFIX",
        source_scope="ALL OFFICIAL SOURCES",
        raw_sku_or_pattern=".3",
        canonical_sku_or_replacement=".30",
        reason="Known suffix rule",
        effective_from=date(2024, 1, 1),
        effective_to=None,
        approved_by="KT",
        approved_date=date(2026, 8, 1),
        evidence="unit test",
        notes="",
    )


def registry(*exceptions: ExceptionRule) -> ManualCheckRegistry:
    return ManualCheckRegistry(exceptions, (preorder_rule(),), (suffix_mapping(),))


def test_approved_exact_match_is_suppressed() -> None:
    cases = [{"record_key": "DH-0001", "issue_type": "CRM missing revenue", "status": "open"}]
    active, suppressed = partition_mismatches(cases, registry(exception_rule()), AS_OF)
    assert active == []
    assert suppressed[0]["manual_check_id"] == "EXC-TEST"
    assert suppressed[0]["status"] == "ignored"


def test_open_or_expired_exception_never_suppresses() -> None:
    cases = [{"record_key": "DH-0001", "issue_type": "CRM MISSING REVENUE"}]
    open_rule = exception_rule(status="OPEN", action="KEEP AS MISMATCH")
    expired_rule = exception_rule(effective_to=date(2025, 12, 31))
    assert partition_mismatches(cases, registry(open_rule), AS_OF)[0] == cases
    assert partition_mismatches(cases, registry(expired_rule), AS_OF)[0] == cases


def test_permanent_preorder_exclusion_ignores_snapshot_quantity_and_value() -> None:
    exact = {"ĐH": "DH-0002", "PRODUCT ID": "SKU.3", "QUANTITY SOLD": 2, "NET REV SOLD": 100}
    changed_quantity = {**exact, "QUANTITY SOLD": 3}
    source = [exact.copy(), changed_quantity.copy()]
    kept, excluded = exclude_preorders(source, registry(), AS_OF)
    assert len(excluded) == 2
    assert excluded[0]["manual_check_id"] == "PREX-TEST"
    assert kept == []
    assert source == [exact, changed_quantity]  # Manual Check annotates copies; it never mutates official input.


def test_order_exclusion_removes_all_rows_for_approved_order() -> None:
    source = [
        {"Order ID": "DH-CANCELLED", "value": 100},
        {"Order ID": "DH-KEEP", "value": 200},
    ]
    base = registry()
    with_orders = ManualCheckRegistry(base.exceptions, base.preorder_exclusions, base.sku_mappings, (order_rule(),))
    kept, excluded = exclude_orders(source, with_orders, AS_OF)
    assert kept == [{"Order ID": "DH-KEEP", "value": 200}]
    assert excluded[0]["manual_check_id"] == "OREX-TEST"


def test_new_unmatched_issue_stays_visible() -> None:
    case = {"record_key": "DH-NEW", "issue_type": "NEW ISSUE"}
    active, suppressed = partition_mismatches([case], registry(exception_rule()), AS_OF)
    assert active == [case]
    assert suppressed == []


def test_migration_snapshot_has_expected_counts() -> None:
    workbook = Path(os.environ.get(
        "MANUAL_CHECK_MIGRATION_WORKBOOK",
        str(Path(__file__).resolve().parents[1] / "input" / "PSI_Manual_Check.xlsx"),
    ))
    if not workbook.exists():
        pytest.skip("Private migration snapshot is unavailable; set MANUAL_CHECK_MIGRATION_WORKBOOK")
    snapshot = workbook.read_bytes()
    expected_hash = "21b0fae9a8ca06b1a89076a31ebd1fd5737ce8bb7b2742d0f1ffcd449002014a"
    if sha256(snapshot).hexdigest() != expected_hash:
        pytest.skip(f"Private migration snapshot SHA-256 differs from {expected_hash}")
    loaded = load_manual_check(snapshot)
    assert loaded.summary(AS_OF) == {
        "exceptions": 110,
        "approved_active_exceptions": 97,
        "open_exceptions": 12,
        "approved_active_preorder_exclusions": 374,
        "approved_active_order_exclusions": 4,
        "approved_active_sku_mappings": 16,
    }
    assert loaded.map_sku("USMUS-11219.3", as_of=AS_OF) == "USMUS-11219.30"


@pytest.mark.parametrize("as_of", [AS_OF, date(2026, 8, 14), date(2026, 8, 27)])
def test_current_workbook_validates_and_filters_as_of(as_of: date) -> None:
    workbook = Path(__file__).resolve().parents[1] / "input" / "PSI_Manual_Check.xlsx"
    if not workbook.exists():
        pytest.skip("Current private input/PSI_Manual_Check.xlsx is not tracked in the repository")
    contents = workbook.read_bytes()
    loaded = load_manual_check(contents)
    # Count directly from workbook cells, independently of registry is_active().
    source = load_workbook(workbook, read_only=True, data_only=True)
    expected = {}
    try:
        for sheet, metric in (
            ("Exceptions", "approved_active_exceptions"),
            ("Preorder Exclusions", "approved_active_preorder_exclusions"),
            ("Order Exclusions", "approved_active_order_exclusions"),
            ("SKU Mappings", "approved_active_sku_mappings"),
        ):
            rows = source[sheet].iter_rows(values_only=True)
            headers = next(
                row for row in rows
                if {"Status", "Effective From", "Effective To"}.issubset(row)
            )
            records = [dict(zip(headers, row)) for row in rows if row[0]]
            assert records, f"{sheet} must exercise the loader"
            expected[metric] = sum(
                str(row["Status"]).strip().upper() == "APPROVED"
                and (not row["Effective From"] or str(row["Effective From"])[:10] <= as_of.isoformat())
                and (not row["Effective To"] or as_of.isoformat() <= str(row["Effective To"])[:10])
                for row in records
            )
            if sheet == "Exceptions":
                expected["exceptions"] = len(records)
                expected["open_exceptions"] = sum(
                    str(row["Status"]).strip().upper() == "OPEN" for row in records
                )
    finally:
        source.close()
    assert loaded.summary(as_of) == expected
