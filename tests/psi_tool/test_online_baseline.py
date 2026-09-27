"""An invalid baseline must not silently suppress the NEW mismatch state."""

import io
from datetime import date

import pytest
from openpyxl import Workbook

from psi_tool.online_baseline import MISMATCH_HEADERS, SHEET_ORDER, verify_baseline


def baseline(title="PSI SUMMARY — 27.08.2026", status="PASS", as_of=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_ORDER:
        workbook.create_sheet(name)
    workbook["PSI Summary"]["A1"] = title
    if as_of is not None:
        workbook["PSI Summary"]["A4"] = "As of"
        workbook["PSI Summary"]["B4"] = as_of
    for column, header in enumerate(MISMATCH_HEADERS, 1):
        workbook["Mismatch"].cell(3, column, header)
    workbook["Checks"]["D4"] = status
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_prior_final_date_comes_from_content():
    assert verify_baseline(
        baseline(), "PSI_Final_27.08.2026.xlsx", date(2026, 9, 3)
    ) == date(2026, 8, 27)


def test_historical_final_with_as_of_field_is_accepted():
    assert verify_baseline(
        baseline("PSI SUMMARY", as_of="2026-08-19"),
        "PSI_Final_19.08.2026.xlsx",
        date(2026, 9, 3),
    ) == date(2026, 8, 19)


def test_conflicting_title_and_as_of_field_are_rejected():
    with pytest.raises(ValueError, match="BASELINE_DATE_MISMATCH"):
        verify_baseline(
            baseline(as_of="2026-08-19"), "PSI Final.xlsx", date(2026, 9, 3)
        )


@pytest.mark.parametrize(
    "title,name,cutoff,status,code",
    [
        (
            "PSI SUMMARY — 03.09.2026",
            "PSI Final.xlsx",
            date(2026, 9, 3),
            "PASS",
            "BASELINE_NOT_EARLIER",
        ),
        (
            "PSI SUMMARY — 04.09.2026",
            "PSI Final.xlsx",
            date(2026, 9, 3),
            "PASS",
            "BASELINE_NOT_EARLIER",
        ),
        (
            "PSI SUMMARY — 27.08.2026",
            "PSI Draft.xlsx",
            date(2026, 9, 3),
            "PASS",
            "BASELINE_FINAL_REQUIRED",
        ),
        (
            "PSI SUMMARY — 27.08.2026",
            "PSI_Final_19.08.2026.xlsx",
            date(2026, 9, 3),
            "PASS",
            "BASELINE_DATE_MISMATCH",
        ),
        (
            "PSI SUMMARY",
            "PSI Final.xlsx",
            date(2026, 9, 3),
            "PASS",
            "BASELINE_DATE_MISSING",
        ),
        (
            "PSI SUMMARY — 27.08.2026",
            "PSI Final.xlsx",
            date(2026, 9, 3),
            "FAIL",
            "BASELINE_CHECKS_FAILED",
        ),
    ],
)
def test_unusable_baseline_rejected(title, name, cutoff, status, code):
    with pytest.raises(ValueError, match=code):
        verify_baseline(baseline(title, status), name, cutoff)
