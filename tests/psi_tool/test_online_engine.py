from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from psi_tool import online_engine
from psi_tool._online_manual_check import (
    EXCEPTION_HEADERS,
    MAPPING_HEADERS,
    ORDER_EXCLUSION_HEADERS,
    PREORDER_HEADERS,
    REQUIRED_SHEETS,
    load_manual_check,
    make_fingerprint,
    make_order_exclusion_fingerprint,
    make_preorder_fingerprint,
)
from psi_tool.online_engine import prepare_payload, validate_payload

AS_OF = date(2026, 9, 3)


@pytest.mark.parametrize(
    "sheet,field",
    [
        ("Preorder Exclusions", "Effective To"),
        ("Order Exclusions", "Effective To"),
        ("Preorder Exclusions", "Evidence"),
        ("Preorder Exclusions", "KT Note"),
    ],
)
def test_permanent_exclusions_cannot_expire_or_lack_evidence(tmp_path, sheet, field):
    path = tmp_path / "manual.xlsx"
    manual_book(path)
    book = load_workbook(path)
    worksheet = book[sheet]
    headers = [cell.value for cell in worksheet[1]]
    cell = worksheet.cell(2, headers.index(field) + 1)
    cell.value = date(2026, 8, 31) if field == "Effective To" else None
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="permanent|reason and evidence"):
        load_manual_check(path)


def write_book(path, sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    book.save(path)
    book.close()


def manual_book(path):
    sheets = {name: [] for name in sorted(REQUIRED_SHEETS)}
    sheets["Exceptions"] = [EXCEPTION_HEADERS]
    sheets["Preorder Exclusions"] = [PREORDER_HEADERS]
    common = {
        "Status": "APPROVED",
        "Effective From": date(2024, 1, 1),
        "Approved By": "TEST",
        "Approved Date": date(2026, 8, 1),
        "Evidence": "synthetic test",
        "Reason": "synthetic test",
        "KT Note": "synthetic fulfilled order",
    }
    for index in (1, 2):
        rule = {
            **common,
            "Exclusion ID": f"PRE-{index}",
            "Order ID": "DH-EXCLUDE",
            "Raw SKU": "SKU",
            "Canonical SKU": "SKU",
            "Action": "EXCLUDE FROM PREORDER",
            "Disposition": "PERMANENT / DONE",
            "Quantity": 999,
            "Net Value": 999,
            "Fingerprint": make_preorder_fingerprint(
                action="EXCLUDE FROM PREORDER",
                order_id="DH-EXCLUDE",
                canonical_sku="SKU",
            ),
        }
        sheets["Preorder Exclusions"].append([rule.get(h) for h in PREORDER_HEADERS])
    order = {
        **common,
        "Exclusion ID": "ORDER-1",
        "Order ID": "DH-CANCELLED",
        "Action": "EXCLUDE ORDER FROM PSI",
        "Scope": "ALL PSI BUSINESS SHEETS",
        "Disposition": "PERMANENT / DONE",
        "Fingerprint": make_order_exclusion_fingerprint(
            action="EXCLUDE ORDER FROM PSI",
            scope="ALL PSI BUSINESS SHEETS",
            order_id="DH-CANCELLED",
        ),
    }
    sheets["Order Exclusions"] = [
        ORDER_EXCLUSION_HEADERS,
        [order.get(h) for h in ORDER_EXCLUSION_HEADERS],
    ]
    mapping = {
        **common,
        "Mapping ID": "MAP-1",
        "Action": "MAP SKU",
        "Match Type": "EXACT",
        "Source Scope": "INVENTORY",
        "Raw SKU / Pattern": "INV-SKU",
        "Canonical SKU / Replacement": "SKU",
        "Fingerprint": make_fingerprint(
            action="MAP SKU",
            scope="SKU",
            raw_sku="INV-SKU",
            canonical_sku="SKU",
            issue_type="EXACT:INVENTORY",
        ),
    }
    sheets["SKU Mappings"] = [
        MAPPING_HEADERS,
        [mapping.get(h) for h in MAPPING_HEADERS],
    ]
    write_book(path, sheets)


@pytest.fixture
def inputs(tmp_path):
    sources = {
        name: tmp_path / f"input-{i}.xlsx"
        for i, name in enumerate(online_engine.SOURCE_LABELS)
    }
    write_book(
        sources["Product Master"],
        {
            "Danh sách": [
                [
                    "Mã hàng hóa",
                    "Tên hàng hóa",
                    "Nguồn gốc",
                    "Mã hãng",
                    "Category",
                    "Sub Category",
                    "Loại hàng hóa",
                ],
                ["SKU", "Product", "Brand", "BC", "Category", "Sub", "Supplier"],
            ]
        },
    )
    orders = ["DH-OPEN", "DH-EXCLUDE", "DH-CANCELLED", "DH-ZERO", "DH-OVER"]
    write_book(
        sources["CRM Sales Order"],
        {
            "Danh sách": [
                [
                    "Số đơn hàng",
                    "Ngày duyệt",
                    "Trạng thái phê duyệt",
                    "Tình trạng",
                    "Khách hàng",
                    "Giá trị đơn hàng",
                ]
            ]
            + [
                [order, AS_OF, "Đã duyệt", "Active", "Customer", 100]
                for order in orders
            ]
            + [["DH-FUTURE", date(2026, 9, 4), "Đã duyệt", "Active", "Customer", 100]],
            "Bảng hàng hóa": [
                ["Số đơn hàng", "Mã hàng hóa", "Số lượng", "Thành tiền sau CK"]
            ]
            + [[order, "SKU", 10, 100] for order in orders]
            + [["DH-OPEN", "SKU", 2, 20]],
        },
    )
    rev_headers = [
        "Ngày hạch toán",
        "Mã hàng",
        "Tổng số lượng bán",
        "Doanh số bán",
        "Chiết khấu",
        "Giá trị trả lại",
        "Giá trị giảm giá",
        "Giá vốn",
        "Đơn hàng",
        "TK Nợ",
        "TK Có",
        "Mã khách hàng",
    ]
    write_book(
        sources["Revenue"],
        {
            "SỔ CHI TIẾT BÁN HÀNG": [
                [],
                [],
                [],
                rev_headers,
                [AS_OF, "SKU", 3, 40, 5, 3, 2, 10, "DH-OPEN", "131", "5111", "KH-OPEN"],
                [
                    AS_OF,
                    "SKU",
                    10,
                    90,
                    0,
                    0,
                    0,
                    20,
                    "DH-ZERO",
                    "131",
                    "5111",
                    "KH-ZERO",
                ],
                [
                    AS_OF,
                    "SKU",
                    11,
                    110,
                    0,
                    0,
                    0,
                    30,
                    "DH-OVER",
                    "131",
                    "5111",
                    "KH-OVER",
                ],
                [
                    AS_OF,
                    "SKU",
                    1,
                    10,
                    0,
                    0,
                    0,
                    2,
                    "DH-CANCELLED",
                    "131",
                    "5111",
                    "KH-CANCELLED",
                ],
                [
                    AS_OF,
                    "SKU",
                    999,
                    999,
                    0,
                    0,
                    0,
                    0,
                    "DH-OPEN",
                    "131",
                    "9999",
                    "KH-OPEN",
                ],
            ]
        },
    )
    inv_headers = ["Tên kho", "Mã kho", "Mã hàng"] + [None] * 8 + ["Cuối kỳ", None]
    write_book(
        sources["Inventory"],
        {
            "TỔNG HỢP TỒN KHO": [
                [],
                [],
                [],
                inv_headers,
                [None] * 11 + ["Số lượng", "Giá trị"],
                ["Main", "A", "INV-SKU"] + [None] * 8 + [7, 70],
                ["Other", "B", "SKU"] + [None] * 8 + [4, 40],
                ["Main", "A", "SKU"] + [None] * 8 + [-3, -30],
                ["Kho Chị Kathy", "K", "SKU"] + [None] * 8 + [99, 99],
            ]
        },
    )
    po = [None] * 78
    for index, val in {
        1: "Pending",
        4: "PO-1",
        18: "SKU",
        19: "Product",
        20: 5,
        26: "F.O.C",
        27: 123.45,
        43: date(2026, 9, 20),
        54: AS_OF,
        76: 7,
        77: 35,
    }.items():
        po[index] = val
    blank = list(po)
    blank[18] = None
    blank[19] = "Unmapped PO"
    done = list(po)
    done[1] = "Done"
    done[20] = 2
    write_book(sources["Purchase/PO"], {"LDL": [[], [], [], [], po, blank, done]})
    write_book(sources["Target"], {"Target": [["Month", "Amount"], [AS_OF, 1000]]})
    manual_book(sources["Manual Check"])
    prior = tmp_path / "prior.xlsx"
    write_book(
        prior,
        {
            "Mismatch": [
                [],
                [],
                [],
                [
                    "OPEN",
                    "CRM / Revenue",
                    "DH-OVER / SKU",
                    "Revenue exceeds approved CRM line",
                    "",
                    "old evidence",
                ],
            ]
        },
    )
    return sources, prior


def test_portable_governed_balances_and_all_relations(inputs, monkeypatch, tmp_path):
    sources, prior = inputs
    monkeypatch.chdir(tmp_path)
    payload = prepare_payload(sources, AS_OF, prior)
    assert sum(isinstance(v, list) for v in payload.values()) == 17
    assert payload["purchase_rows"][0][13] == 123.45
    assert payload["purchase_rows"][0][14] == "2026-09-20"
    assert payload["summary"]["inventory_qty"] == 11
    assert len(payload["inventory_rows"]) == 2
    assert payload["revenue_rows"][0][9] == 30
    assert payload["revenue_rows"][0][12] == "KH-OPEN"
    assert payload["preorder_rows"][0][8:10] == [9, 90]
    assert payload["preorder_excluded_rows"][0][8:10] == [10, 100]
    assert len(payload["preorder_excluded_rows"]) == 1
    assert len(payload["po_excluded_rows"]) == 1
    assert payload["psi_product_rows"][0][14] == 5
    assert payload["summary"]["new_mismatch_rows"] == 1
    assert payload["gate_evidence"]["active_preorder_exclusion_rows"] == 2
    assert payload["gate_evidence"]["active_preorder_exclusion_unique_keys"] == 1
    assert all(row[0] != "DH-CANCELLED" for row in payload["crm_final_rows"])
    report = validate_payload(payload, sources, AS_OF, prior)
    assert report["status"] == "PASS", report
    assert report["failure_count"] == 0
    assert report["warnings"]


@pytest.mark.parametrize(
    "mutation",
    [
        "revenue",
        "warehouse",
        "duplicate",
        "target",
        "category",
        "gap_alias",
        "new_mismatch",
        "hash",
        "baseline",
        "shape",
        "nan",
        "crm_customer",
        "margin",
        "gates",
        "purchase_cost",
        "crm_product_date",
        "removed_issue",
    ],
)
def test_independent_validator_rejects_tampering(inputs, mutation):
    sources, prior = inputs
    payload = prepare_payload(sources, AS_OF, prior)
    if mutation == "revenue":
        payload["revenue_rows"][0][9] += 1
    elif mutation == "warehouse":
        payload["inventory_rows"][0][10] = "WRONG"
    elif mutation == "duplicate":
        payload["preorder_excluded_rows"].append(
            deepcopy(payload["preorder_excluded_rows"][0])
        )
    elif mutation == "target":
        payload["target_rows"][1][1] += 1
    elif mutation == "category":
        payload["category_rows"] = []
    elif mutation == "gap_alias":
        payload["data_gaps_rows"] = []
    elif mutation == "new_mismatch":
        payload["new_mismatch_keys"] = []
    elif mutation == "hash":
        payload["source_hashes"].pop("Revenue")
    elif mutation == "baseline":
        payload["prior_psi_sha256"] = ""
    elif mutation == "shape":
        payload["revenue_rows"] = [[]]
    elif mutation == "nan":
        payload["psi_product_rows"][0][7] = float("nan")
    elif mutation == "crm_customer":
        payload["crm_final_rows"][0][1] = "WRONG"
    elif mutation == "purchase_cost":
        payload["purchase_rows"][0][11] += 1
    elif mutation == "crm_product_date":
        payload["crm_product_rows"][0][5] = "2020-01-01"
    elif mutation == "gates":
        payload["gates"] = [{"check": "fake", "status": "PASS"}]
    elif mutation == "removed_issue":
        payload["mismatch_rows"] = []
        payload["new_mismatch_keys"] = []
        payload["summary"]["mismatch_rows"] = 0
        payload["summary"]["new_mismatch_rows"] = 0
    elif mutation == "margin":
        payload["psi_product_rows"][0][10] = 123
    report = validate_payload(payload, sources, AS_OF, prior)
    assert report["status"] == "FAIL", mutation
    assert report["failure_count"] > 0


def test_explicit_baseline_and_source_labels_required(inputs):
    sources, prior = inputs
    with pytest.raises(ValueError, match="baseline"):
        prepare_payload(sources, AS_OF, None)
    with pytest.raises(ValueError, match="seven"):
        prepare_payload({"Revenue": sources["Revenue"]}, AS_OF, prior)


def test_concurrent_runs_do_not_share_dates_or_state(inputs):
    sources, prior = inputs
    with ThreadPoolExecutor(max_workers=2) as executor:
        future = executor.submit(prepare_payload, sources, AS_OF, prior)
        early = executor.submit(prepare_payload, sources, date(2026, 9, 2), prior)
    assert future.result()["summary"]["revenue_rows"] == 3
    assert early.result()["summary"]["revenue_rows"] == 0
    assert early.result()["as_of"] == "2026-09-02"


def test_inputs_must_not_change_during_preparation(inputs, monkeypatch):
    sources, prior = inputs
    actual = online_engine.prepare

    def mutate(*args):
        result = actual(*args)
        write_book(sources["Target"], {"Target": [["Changed"]]})
        return result

    monkeypatch.setattr(online_engine, "prepare", mutate)
    with pytest.raises(ValueError, match="changed"):
        prepare_payload(sources, AS_OF, prior)


def test_blank_order_mismatch_identity_is_trimmed(inputs):
    sources, prior = inputs
    book = load_workbook(sources["Revenue"])
    book["SỔ CHI TIẾT BÁN HÀNG"].append(
        [AS_OF, "SKU", 1, 0, 0, 0, 0, 20, None, "131", "5111"]
    )
    book.save(sources["Revenue"])
    book.close()
    payload = prepare_payload(sources, AS_OF, prior)
    assert any(row[2] == "/ SKU" for row in payload["mismatch_rows"])
    assert validate_payload(payload, sources, AS_OF, prior)["status"] == "PASS"


def test_crm_specific_mapping_is_recomputed_in_the_same_scope(inputs):
    sources, prior = inputs
    book = load_workbook(sources["Manual Check"])
    rule = {
        "Mapping ID": "MAP-CRM",
        "Action": "MAP SKU",
        "Match Type": "EXACT",
        "Status": "APPROVED",
        "Source Scope": "CRM",
        "Raw SKU / Pattern": "CRM-SKU",
        "Canonical SKU / Replacement": "SKU",
        "Approved By": "TEST",
        "Approved Date": AS_OF,
        "Fingerprint": make_fingerprint(
            action="MAP SKU",
            scope="SKU",
            raw_sku="CRM-SKU",
            canonical_sku="SKU",
            issue_type="EXACT:CRM",
        ),
    }
    book["SKU Mappings"].append([rule.get(h) for h in MAPPING_HEADERS])
    book.save(sources["Manual Check"])
    book.close()
    book = load_workbook(sources["CRM Sales Order"])
    book["Bảng hàng hóa"].cell(2, 2, "CRM-SKU")
    book.save(sources["CRM Sales Order"])
    book.close()
    payload = prepare_payload(sources, AS_OF, prior)
    assert all(row[1] == "SKU" for row in payload["crm_product_rows"])
    assert validate_payload(payload, sources, AS_OF, prior)["status"] == "PASS"
