"""Portable export parity and hostile workbook regression tests."""

# Synthetic JSON deliberately exercises malformed dynamic inputs.
# pyright: reportAny=false, reportExplicitAny=false, reportUnknownMemberType=false
# pyright: reportUnusedCallResult=false, reportImplicitStringConcatenation=false
import io
import json
import math
import os
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from psi_tool._online_workbook_schema import SCHEMA
from psi_tool.online_workbook import build_workbook, validate_workbook


@pytest.fixture
def payload() -> dict[str, Any]:
    result: dict[str, Any] = {
        "as_of": "2026-09-03",
        "period_start": "2024-01-01",
        "sources": {"Revenue": "/private/revenue.xlsx"},
        "source_hashes": {"Revenue": "a" * 64},
        "prior_psi_final": "/private/PSI_Final_27.08.2026.xlsx",
        "prior_psi_sha256": "b" * 64,
        "gates": [
            {
                "check": "independent_validation",
                "status": "PASS",
                "expected": True,
                "actual": True,
            }
        ],
        "mismatch_rows": [
            [
                "OPEN",
                "CRM",
                "P01",
                "SKU not found",
                "Review",
                '=HYPERLINK("https://invalid.example","x")',
            ]
        ],
        "new_mismatch_keys": [["CRM", "P01", "SKU not found"]],
        "target_rows": [
            ["NO.", "BRAND", "CODE", "ROOM", "T1", "T2", "T3", "C1", "C2", "IGNORED"],
            [1, "Sample", "S", "Main", 1, 2, 3, 4, 5, "=1+1"],
        ],
    }
    for spec in SCHEMA:
        if spec["key"]:
            row: list[Any] = [""] * len(spec["headers"])
            row[0] = "P01"
            for col in spec["quantity"] + spec["money"] + spec["percent"]:
                row[col] = 2
            result[spec["key"]] = [row]
    result["psi_product_rows"][0][7:9] = [100, 40]
    result["revenue_rows"][0][8:11] = [2, 100, 40]
    result["brand_rows"] = [
        ["=1+1", "https://example.com"],
        ["+1+1", "@SUM(A1)"],
        ["-1+1", "001"],
    ]
    return result


def _mutate(content: bytes, entry: str, change: Callable[[bytes], bytes]) -> bytes:
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            value = source.read(item.filename)
            target.writestr(item, change(value) if item.filename == entry else value)
    return out.getvalue()


def test_synthetic_export_is_valid_deterministic_and_safe(
    payload: dict[str, Any],
) -> None:
    content = build_workbook(payload)
    assert content == build_workbook(payload)
    report = validate_workbook(content, payload)
    assert report["status"] == "PASS", report
    assert report["formula_count"] == 13
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert len(wb.sheetnames) == 17
    assert wb["Brand"]["A4"].value == "=1+1"
    assert wb["Brand"]["A4"].data_type == "s"
    assert wb["Brand"]["B4"].hyperlink is None
    assert wb["PSI by Product"]["J4"].value == "=H4-I4"
    assert wb["PSI Summary"]["B5"].value == "=SUM('Revenue'!J4:J4)"
    cached = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    assert cached["PSI by Product"]["J4"].value == 60
    assert cached["PSI by Product"]["K4"].value == 0.6
    assert cached["PSI Summary"]["B10"].value == 1
    assert wb["Target"].max_column == 9
    assert all(s.freeze_panes == "A4" for s in wb)


@pytest.mark.parametrize(
    "gates", [None, [], [{"status": "FAIL"}], [{"status": "UNKNOWN"}]]
)
def test_refuses_absent_or_failed_gates(payload: dict[str, Any], gates: object) -> None:
    payload["gates"] = gates
    with pytest.raises(ValueError, match="gate"):
        build_workbook(payload)


def test_refuses_missing_hash_or_bad_width(payload: dict[str, Any]) -> None:
    payload["source_hashes"] = {}
    with pytest.raises(ValueError, match="SHA-256"):
        build_workbook(payload)
    payload["source_hashes"] = {"Revenue": "a" * 64}
    payload["revenue_rows"][0].append("extra")
    with pytest.raises(ValueError, match="width"):
        build_workbook(payload)


@pytest.mark.parametrize(
    ("entry", "old", "new"),
    [
        ("xl/worksheets/sheet7.xml", b"<f>H4-I4</f>", b"<f>H4+I4</f>"),
        ("xl/worksheets/sheet7.xml", b"<v>60</v>", b"<v>999</v>"),
        (
            "xl/worksheets/sheet7.xml",
            b"<f>H4-I4</f><v>60</v>",
            b"<f>H4-I4</f><v>#REF!</v>",
        ),
        ("xl/worksheets/sheet4.xml", b'sqref="A4:G4"', b'sqref="G4"'),
        ("xl/worksheets/sheet1.xml", b'topLeftCell="A4"', b'topLeftCell="A2"'),
        ("xl/worksheets/sheet1.xml", b'width="31"', b'width="32"'),
        ("xl/sharedStrings.xml", b"a" * 64, b"c" * 64),
    ],
)
def test_independent_validator_detects_tampering(
    payload: dict[str, Any], entry: str, old: bytes, new: bytes
) -> None:
    original = build_workbook(payload)

    def change(value: bytes) -> bytes:
        assert old in value
        return value.replace(old, new, 1)

    report = validate_workbook(_mutate(original, entry, change), payload)
    assert report["status"] == "FAIL", report
    assert report["failure_count"] >= 1


def test_validator_rejects_external_links_and_injected_formulas(
    payload: dict[str, Any],
) -> None:
    content = build_workbook(payload)
    injected = _mutate(
        content,
        "xl/worksheets/sheet8.xml",
        lambda value: value.replace(b'<c r="A4"', b'<c r="A4"', 1).replace(
            b"</row>",
            b'<c r="Z1"><f>WEBSERVICE("https://invalid.example")</f><v>0</v></c></row>',
            1,
        ),
    )
    assert validate_workbook(injected, payload)["status"] == "FAIL"
    external = _mutate(
        content,
        "xl/_rels/workbook.xml.rels",
        lambda value: value.replace(
            b"</Relationships>",
            b'<Relationship Id="rId99" Type="external" '
            b'Target="https://invalid.example" TargetMode="External"/></Relationships>',
        ),
    )
    assert validate_workbook(external, payload)["status"] == "FAIL"


def test_empty_data_and_invalid_archive(payload: dict[str, Any]) -> None:
    for key in list(payload):
        if key.endswith("_rows") or key == "new_mismatch_keys":
            payload[key] = []
    report = validate_workbook(build_workbook(payload), payload)
    assert report["status"] == "PASS", report
    assert report["formula_count"] == 11
    assert validate_workbook(b"not an XLSX", payload)["status"] == "FAIL"


def test_golden_offline_parity_when_fixture_is_available() -> None:
    """Opt-in private fixture; no private source data is committed to tests."""
    directory = os.environ.get("PSI_GOLDEN_DIR")
    if not directory:
        pytest.skip("Set PSI_GOLDEN_DIR to the verified 03/09 offline output directory")
    root = Path(directory)
    payload = json.loads((root / "psi_20260903_data.json").read_text())
    content = build_workbook(payload)
    report = validate_workbook(content, payload)
    assert report["status"] == "PASS", report
    assert report["formula_count"] == 8565
    with (
        zipfile.ZipFile(root / "PSI_Final_03.09.2026.xlsx") as offline,
        zipfile.ZipFile(io.BytesIO(content)) as online,
    ):
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for i in range(1, 18):
            first, second = (
                ET.fromstring(z.read(f"xl/worksheets/sheet{i}.xml"))  # noqa: S314 - trusted local golden fixture
                for z in (offline, online)
            )
            for tag in ["mergeCells/mergeCell"]:
                query = "/".join("s:" + part for part in tag.split("/"))
                assert {x.get("ref") for x in first.findall(query, ns)} == {
                    x.get("ref") for x in second.findall(query, ns)
                }

            def widths(xml: ET.Element) -> dict[int, float]:
                return {
                    i: float(c.attrib["width"])
                    for c in xml.findall("s:cols/s:col", ns)
                    for i in range(int(c.attrib["min"]), int(c.attrib["max"]) + 1)
                }

            assert widths(first) == widths(second)
    expected = openpyxl.load_workbook(
        root / "PSI_Final_03.09.2026.xlsx", read_only=True
    )
    actual = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    assert expected.sheetnames == actual.sheetnames
    for before, after in zip(expected, actual, strict=True):
        cast("Any", before).calculate_dimension(force=True)
        assert (before.max_row, before.max_column) == (after.max_row, after.max_column)
        for left, right in zip(before.iter_rows(), after.iter_rows(), strict=True):
            for a, b in zip(left, right, strict=True):
                x, y = a.value or "", b.value or ""
                if x or y:
                    assert a.number_format == b.number_format, (
                        before.title,
                        a.coordinate,
                        "number_format",
                        a.number_format,
                        b.number_format,
                    )
                    assert (a.font.name, a.font.sz) == (b.font.name, b.font.sz), (
                        before.title,
                        a.coordinate,
                        "font",
                        a.font.name,
                        a.font.sz,
                        b.font.name,
                        b.font.sz,
                    )
                if isinstance(x, (int, float)) and isinstance(y, (int, float)):
                    assert math.isclose(x, y, rel_tol=1e-14, abs_tol=1e-7), (
                        before.title,
                        a.coordinate,
                    )
                else:
                    assert x == y, (before.title, a.coordinate)
                assert (a.data_type == "f") == (b.data_type == "f"), (
                    before.title,
                    a.coordinate,
                )
    expected.close()
    actual.close()


def test_formula_inputs_cannot_hide_excel_errors(payload: dict[str, Any]) -> None:
    payload["psi_product_rows"][0][7] = "not a number"
    with pytest.raises(ValueError, match="numeric"):
        build_workbook(payload)


def test_validator_checks_numeric_formats_and_header_style(
    payload: dict[str, Any],
) -> None:
    content = build_workbook(payload)
    tampered = _mutate(
        content,
        "xl/styles.xml",
        lambda value: value.replace(b"#,##0.0;[Red](#,##0.0);-", b"0.00", 1),
    )
    assert validate_workbook(tampered, payload)["status"] == "FAIL"
    tampered = _mutate(
        content, "xl/styles.xml", lambda value: value.replace(b"FF17365D", b"FF000000")
    )
    assert validate_workbook(tampered, payload)["status"] == "FAIL"


def test_validator_requires_new_highlight_color(payload: dict[str, Any]) -> None:
    content = build_workbook(payload)
    tampered = _mutate(
        content, "xl/styles.xml", lambda value: value.replace(b"FFFCE4D6", b"FF000000")
    )
    assert validate_workbook(tampered, payload)["status"] == "FAIL"
