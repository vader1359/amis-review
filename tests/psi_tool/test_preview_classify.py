"""Routing is content-based and never silently chooses ambiguous exports."""

import io
import unicodedata

import pytest
from openpyxl import Workbook
from web.preview_classify import classify_files, validate_reusable


def workbook(sheets: dict[str, tuple[int, list[str]]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for name, (row, headers) in sheets.items():
        sheet = book.create_sheet(name)
        for column, value in enumerate(headers, start=1):
            sheet.cell(row, column, value)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


PRODUCT = {
    "Danh sách": (
        1,
        [
            "Mã hàng hóa",
            "Tên hàng hóa",
            "Nguồn gốc",
            "Category",
            "Sub Category",
        ],
    )
}
CRM = {
    "Danh sách": (
        1,
        [
            "Số đơn hàng",
            "Ngày duyệt",
            "Trạng thái phê duyệt",
            "Tình trạng",
        ],
    ),
    "Bảng hàng hóa": (
        1,
        [
            "Số đơn hàng",
            "Mã hàng hóa",
            "Số lượng",
            "Tổng tiền",
        ],
    ),
}
REVENUE = {
    "SỔ CHI TIẾT BÁN HÀNG": (
        4,
        [
            "Ngày hạch toán",
            "Mã hàng",
            "Tổng số lượng bán",
            "Doanh số bán",
            "Chiết khấu",
            "Giá trị trả lại",
            "Giá trị giảm giá",
            "Giá vốn",
            "Đơn hàng",
            "TK Nợ",
            "TK Có",
        ],
    )
}
INVENTORY = {"TỔNG HỢP TỒN KHO": (4, ["Tên kho", "Mã kho", "Mã hàng", "Cuối kỳ"])}


def test_four_renamed_files_are_sorted_by_content() -> None:
    inputs = [INVENTORY, PRODUCT, CRM, REVENUE]
    results = classify_files(
        [(f"renamed-{i}.xlsx", workbook(s)) for i, s in enumerate(inputs)]
    )
    assert [item["role"] for item in results] == [
        "inventory",
        "product",
        "crm",
        "revenue",
    ]
    assert [item["index"] for item in results] == list(range(4))


def test_fake_filename_does_not_route_unknown_schema() -> None:
    result = classify_files(
        [("CRM_Product_Master.xlsx", workbook({"Sheet1": (1, ["x"])}))]
    )
    assert result[0]["role"] is None
    assert result[0]["candidates"] == []


def test_duplicate_and_ambiguous_files_are_preserved() -> None:
    duplicate = classify_files(
        [("a.xlsx", workbook(PRODUCT)), ("b.xlsx", workbook(PRODUCT))]
    )
    assert [r["role"] for r in duplicate] == ["product", "product"]
    hybrid = {**CRM, "Danh sách": (1, CRM["Danh sách"][1] + PRODUCT["Danh sách"][1])}
    result = classify_files([("hybrid.xlsx", workbook(hybrid))])[0]
    assert result["role"] is None
    assert set(result["candidates"]) == {"crm", "product"}


def test_missing_crm_items_headers_does_not_guess() -> None:
    result = classify_files(
        [("orders.xlsx", workbook({"Danh sách": CRM["Danh sách"]}))]
    )
    assert result[0]["role"] is None


def test_decomposed_accents_and_whitespace_are_supported() -> None:
    row = [unicodedata.normalize("NFD", f"  {x} ") for x in PRODUCT["Danh sách"][1]]
    assert (
        classify_files([("a.xlsx", workbook({"Danh sách": (1, row)}))])[0]["role"]
        == "product"
    )


def test_headers_must_be_at_engine_read_position() -> None:
    assert (
        classify_files(
            [("a.xlsx", workbook({"Danh sách": (2, PRODUCT["Danh sách"][1])}))]
        )[0]["role"]
        is None
    )


def test_target_reuse_requires_target_headers() -> None:
    valid = workbook(
        {"Target": (2, ["BRAND", "BRAND CODE", "Internal Target value 2027"])}
    )
    validate_reusable("target", "renamed.xlsx", valid)
    with pytest.raises(ValueError, match="SOURCE_SCHEMA_INVALID"):
        validate_reusable("target", "Target.xlsx", workbook(PRODUCT))


def test_purchase_reuse_rejects_shifted_positional_columns() -> None:
    headers = [""] * 78
    for index, label in {
        18: "MÃ MISA",
        19: "TÊN HÀNG HÓA",
        20: "SL",
        76: "Giá nhập kho (chưa VAT)",
        77: "Giá trị nhập kho (theo SL)",
    }.items():
        headers[index] = label
    validate_reusable("purchase", "a.xlsx", workbook({"LDL": (4, headers)}))
    with pytest.raises(ValueError, match="SOURCE_SCHEMA_INVALID"):
        validate_reusable("purchase", "a.xlsx", workbook({"LDL": (4, ["", *headers])}))


def test_corrupt_file_does_not_abort_other_classifications() -> None:
    results = classify_files(
        [("bad.xlsx", b"not-a-workbook"), ("good.xlsx", workbook(PRODUCT))]
    )
    assert results[0]["role"] is None
    assert results[1]["role"] == "product"
