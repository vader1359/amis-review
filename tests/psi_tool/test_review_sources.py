import io
import stat
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from web.review_sources import SourceBundleError, SourceBundleStore, apply_exclusions

from psi_tool._online_manual_check import (
    EXCEPTION_HEADERS,
    MAPPING_HEADERS,
    ORDER_EXCLUSION_HEADERS,
    PREORDER_HEADERS,
    REQUIRED_SHEETS,
    load_manual_check,
)
from psi_tool.online_pipeline import SOURCE_LABELS, SourceSnapshot


def sources():
    book = Workbook()
    book.remove(book.active)
    for name in REQUIRED_SHEETS:
        book.create_sheet(name)
    for name, headers in [
        ("Exceptions", EXCEPTION_HEADERS),
        ("Preorder Exclusions", PREORDER_HEADERS),
        ("Order Exclusions", ORDER_EXCLUSION_HEADERS),
        ("SKU Mappings", MAPPING_HEADERS),
    ]:
        book[name].append(headers)
    output = io.BytesIO()
    book.save(output)
    book.close()
    snapshot = SourceSnapshot.from_bytes("source.xlsx", output.getvalue())
    return dict.fromkeys(SOURCE_LABELS, snapshot), snapshot


def test_bundle_roundtrip_private_immutable(tmp_path):
    store = SourceBundleStore(tmp_path / "bundles")
    selected, prior = sources()
    identifier = "a" * 32
    store.save(identifier, selected, prior)
    assert store.load(identifier) == (selected, prior)
    store.save(identifier, selected, prior)
    assert store.available(identifier)
    assert stat.S_IMODE(store.directory.stat().st_mode) == 0o700
    assert all(
        stat.S_IMODE(p.stat().st_mode) == 0o600
        for p in (store.directory / identifier).iterdir()
    )
    altered = {**selected, "crm": SourceSnapshot.from_bytes("other.xlsx", b"altered")}
    with pytest.raises(SourceBundleError, match="IMMUTABLE"):
        store.save(identifier, altered, prior)
    blob = store.directory / identifier / (prior.sha256 + ".xlsx")
    blob.write_bytes(b"tampered")
    assert not store.available(identifier)
    with pytest.raises(SourceBundleError, match="CHECKSUM"):
        store.load(identifier)


def test_reject_symlinks_and_unsafe_ids(tmp_path):
    store = SourceBundleStore(tmp_path / "bundles")
    selected, prior = sources()
    with pytest.raises(SourceBundleError):
        store.save("../escape", selected, prior)
    (store.directory / ("a" * 32)).symlink_to(tmp_path, target_is_directory=True)
    with pytest.raises(SourceBundleError, match="UNSAFE"):
        store.save("a" * 32, selected, prior)


def test_concurrent_bundle_publication_and_symlink_blob(tmp_path):
    store = SourceBundleStore(tmp_path / "bundles")
    selected, prior = sources()
    identifier = "b" * 32
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [
            pool.submit(store.save, identifier, selected, prior) for _ in range(2)
        ]
        for future in futures:
            future.result()
    assert store.load(identifier) == (selected, prior)
    blob = store.directory / identifier / (prior.sha256 + ".xlsx")
    outside = tmp_path / "outside.xlsx"
    outside.write_bytes(prior.content)
    blob.unlink()
    blob.symlink_to(outside)
    assert not store.available(identifier)


def test_manifest_tampering_fails_closed(tmp_path):
    store = SourceBundleStore(tmp_path / "bundles")
    selected, prior = sources()
    store.save("c" * 32, selected, prior)
    manifest = store.directory / ("c" * 32) / "manifest.json"
    manifest.write_text(manifest.read_text().replace("source.xlsx", "renamed.xlsx"))
    assert not store.available("c" * 32)


def test_governed_exclusions_keep_sources_and_are_idempotent():
    selected, _ = sources()
    cutoff = date(2026, 9, 3)
    proposals = [
        dict(
            id="p1",
            action="exclude_preorder",
            order_id="dh-1",
            sku="sku-1",
            note="=not a formula",
            author="KT",
        ),
        dict(
            id="p2",
            action="exclude_order",
            order_id="dh-2",
            note="Cancelled",
            author="KT",
        ),
    ]
    updated = apply_exclusions(selected, proposals, cutoff, "Approver")
    assert all(updated[k] is selected[k] for k in SOURCE_LABELS if k != "manual_check")
    assert updated["manual_check"].content != selected["manual_check"].content
    registry = load_manual_check(updated["manual_check"].content)
    assert registry.matching_preorder_exclusion(
        order_id="DH-1", sku="SKU-1", quantity=999, net_value=999, as_of=cutoff
    )
    assert registry.matching_order_exclusion("DH-2", cutoff)
    assert not registry.matching_order_exclusion("DH-2", date(2026, 9, 2))
    assert apply_exclusions(updated, proposals, cutoff, "Approver") == updated
    with pytest.raises(SourceBundleError, match="CONFLICT"):
        apply_exclusions(
            updated, [{**proposals[0], "note": "Changed"}], cutoff, "Approver"
        )
    workbook = load_workbook(io.BytesIO(updated["manual_check"].content))
    cell = workbook["Preorder Exclusions"].cell(
        2, PREORDER_HEADERS.index("KT Note") + 1
    )
    assert cell.value == "=not a formula" and cell.data_type == "s"
    workbook.close()


def test_same_exclusion_replay_produces_identical_snapshot_bytes() -> None:
    selected, _ = sources()
    decision = {
        "id": "review-test",
        "action": "exclude_order",
        "order_id": "DH-OPEN",
        "note": "Cancelled",
        "author": "KT",
        "approved_by": "Approver",
        "effective_from": "2026-09-03",
    }
    from web.review_sources import merge_applied_exclusions

    first = merge_applied_exclusions(selected, [decision], date(2026, 9, 3))
    time.sleep(2.1)
    second = merge_applied_exclusions(selected, [decision], date(2026, 9, 3))

    assert first["manual_check"] == second["manual_check"]
    assert selected["manual_check"].content != first["manual_check"].content
    registry = load_manual_check(second["manual_check"].content)
    rule = registry.order_exclusions[0]
    assert rule.approved_by == "Approver"
    assert rule.effective_from == date(2026, 9, 3)


@pytest.mark.parametrize(
    "change", [{"action": "delete"}, {"note": ""}, {"sku": ""}, {"author": ""}]
)
def test_invalid_exclusion_rejected(change):
    selected, _ = sources()
    proposal = dict(
        id="p",
        action="exclude_preorder",
        order_id="DH",
        sku="SKU",
        note="Error",
        author="KT",
    )
    with pytest.raises(SourceBundleError):
        apply_exclusions(
            selected, [{**proposal, **change}], date(2026, 9, 3), "Approver"
        )


def test_permanent_decisions_carry_forward_with_original_approval():
    from web.review_sources import merge_applied_exclusions

    selected, _ = sources()
    decisions = [
        dict(
            id="old",
            action="exclude_order",
            order_id="DH-1",
            note="Cancelled",
            author="KT",
            approved_by="Approver A",
            effective_from="2026-09-03",
        ),
        dict(
            id="later",
            action="exclude_preorder",
            order_id="DH-2",
            sku="SKU",
            note="Error",
            author="KT",
            approved_by="Approver B",
            effective_from="2026-09-10",
        ),
    ]
    first = merge_applied_exclusions(selected, decisions, date(2026, 9, 3))
    registry = load_manual_check(first["manual_check"].content)
    assert len(registry.order_exclusions) == 1
    assert not registry.preorder_exclusions
    second = merge_applied_exclusions(first, decisions, date(2026, 9, 10))
    registry = load_manual_check(second["manual_check"].content)
    rule = registry.order_exclusions[0]
    assert rule.effective_from == rule.approved_date == date(2026, 9, 3)
    assert rule.approved_by == "Approver A"
    assert len(registry.preorder_exclusions) == 1
    assert merge_applied_exclusions(second, decisions, date(2026, 9, 17)) == second
    restored = merge_applied_exclusions(selected, decisions, date(2026, 9, 17))
    assert load_manual_check(restored["manual_check"].content) == registry


@pytest.mark.parametrize(
    "decision",
    [
        {},
        {"effective_from": "invalid", "approved_by": "A"},
        {"effective_from": "2026-09-03", "approved_by": ""},
    ],
)
def test_carry_forward_rejects_invalid_approval_provenance(decision):
    from web.review_sources import merge_applied_exclusions

    selected, _ = sources()
    with pytest.raises(SourceBundleError, match="APPROVAL_INVALID"):
        merge_applied_exclusions(selected, [decision], date(2026, 9, 17))
