"""Immutable Draft snapshots in Neon with optional private host artifact storage.

The migration is applied by an operator, never during a request. Downloads are
read from verified original bytes or reconstructed with the matching renderer.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import platform
import re
import stat
import uuid
import zlib
from contextlib import contextmanager, nullcontext
from datetime import date, datetime, time
from importlib.metadata import version
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

import openpyxl
import psycopg
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb

from psi_tool import _online_workbook_schema, online_workbook
from psi_tool.online_pipeline import DraftResult

MAX_PAYLOAD_BYTES = 100 * 1024 * 1024
MAX_ARTIFACT_BYTES = 100 * 1024 * 1024
REPORT_FIELDS = (
    "id, as_of, created_at, state, input_hash, workbook_sha256, "
    "summary, gates, sheet_count"
)
SHEET_FIELDS = (
    "ordinal, name, row_count, column_count, populated_rows, cell_count, sha256"
)


class ReportStoreError(ValueError):
    """Public, sanitized failure with no driver text or private input values."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def _json(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode()


def _sha(value: Any) -> str:
    return hashlib.sha256(_json(value)).hexdigest()


def renderer_fingerprint() -> str:
    """Pin renderer/schema source and byte-producing dependency/runtime versions."""
    return _sha(
        {
            "format": "psi-preview-renderer-v1",
            "modules": {
                module.__name__: hashlib.sha256(
                    Path(module.__file__).read_bytes()
                ).hexdigest()
                for module in (online_workbook, _online_workbook_schema)
            },
            "xlsxwriter": version("XlsxWriter"),
            "python": platform.python_version(),
            "zlib": zlib.ZLIB_RUNTIME_VERSION,
        }
    )


def _typed(value: Any, data_type: str) -> tuple[str, Any]:
    if isinstance(value, datetime):
        return "datetime", value.isoformat(timespec="microseconds")
    if isinstance(value, date):
        return "date", value.isoformat()
    if isinstance(value, time):
        return "time", value.isoformat(timespec="microseconds")
    if value is None:
        return "blank", None
    if isinstance(value, bool):
        return "boolean", value
    if isinstance(value, (int, float)):
        return "number", value
    return ("error" if data_type == "e" else "text"), value


def snapshot_workbook(content: bytes) -> list[dict[str, Any]]:
    """Capture ordered sparse cells, their Excel positions, formulas and caches."""
    authored = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=False
    )
    cached = None
    try:
        cached = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        expected = [sheet["name"] for sheet in _online_workbook_schema.SCHEMA]
        if authored.sheetnames != expected or cached.sheetnames != expected:
            raise ReportStoreError("REPORT_SHEET_SET_INVALID")
        result = []
        for ordinal, sheet in enumerate(authored, 1):
            rows = []
            cell_count = 0
            for row_number, (cells, values) in enumerate(
                zip(sheet.iter_rows(), cached[sheet.title].iter_rows(), strict=True), 1
            ):
                items = []
                for column, (cell, value) in enumerate(
                    zip(cells, values, strict=True), 1
                ):
                    if cell.value is None:
                        continue
                    kind, encoded = _typed(cell.value, cell.data_type)
                    item = {"column": column, "type": kind, "value": encoded}
                    if cell.data_type == "f":
                        cache_type, cache_value = _typed(value.value, value.data_type)
                        item.update(
                            type="formula",
                            formula=cell.value,
                            cached_type=cache_type,
                            cached_value=cache_value,
                        )
                    items.append(item)
                if items:
                    rows.append({"row_number": row_number, "cells": items})
                    cell_count += len(items)
            metadata = {
                "ordinal": ordinal,
                "name": sheet.title,
                "row_count": sheet.max_row,
                "column_count": sheet.max_column,
                "populated_rows": len(rows),
                "cell_count": cell_count,
            }
            result.append(
                {**metadata, "sha256": _sha({**metadata, "rows": rows}), "rows": rows}
            )
        return result
    finally:
        authored.close()
        if cached is not None:
            cached.close()


def _metadata(row: dict[str, Any]) -> dict[str, Any]:
    return {
        **row,
        "as_of": row["as_of"].isoformat(),
        "created_at": row["created_at"].isoformat(),
        "storage": "neon",
        "validation": {"payload": "PASS", "parquet": "PASS", "workbook": "PASS"},
        "download_url": f"/api/drafts/{row['id']}/download",
    }


def _report_id(value: str) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-f]{32}", value):
        raise ReportStoreError("ONLINE_REPORT_NOT_FOUND")
    return value


def _page(limit: int, offset: int, maximum: int) -> None:
    if (
        type(limit) is not int
        or type(offset) is not int
        or not 1 <= limit <= maximum
        or not 0 <= offset <= 250003
    ):
        raise ReportStoreError("REPORT_PAGE_INVALID")


class OnlineReportStore:
    """Store verified drafts using a private, explicitly configured Neon DSN."""

    def __init__(self, database_url: str, artifact_dir: Path | None = None):
        try:
            parsed = urlsplit(database_url)
            valid = (
                parsed.scheme in {"postgres", "postgresql"}
                and parsed.hostname
                and parsed.hostname.endswith(".neon.tech")
            )
        except (TypeError, ValueError):
            valid = False
        if not valid:
            raise ReportStoreError("ONLINE_DATABASE_CONFIG_INVALID")
        self._database_url = database_url
        configured = artifact_dir if artifact_dir is not None else os.environ.get(
            "PSI_REPORT_ARTIFACT_DIR"
        )
        self._artifact_dir = Path(configured).absolute() if configured else None

    @contextmanager
    def _artifact_directory(self):
        """Open a private directory without following any symlink components."""
        descriptor = None
        try:
            path = self._artifact_dir
            if path is None or ".." in path.parts:
                raise ReportStoreError("REPORT_ARTIFACT_CONFIG_INVALID")
            flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW
            descriptor = os.open(path.anchor, flags)
            for component in path.parts[1:]:
                try:
                    os.mkdir(component, mode=0o700, dir_fd=descriptor)
                except FileExistsError:
                    pass
                child = os.open(component, flags, dir_fd=descriptor)
                os.close(descriptor)
                descriptor = child
            if os.fstat(descriptor).st_mode & 0o077:
                raise ReportStoreError("REPORT_ARTIFACT_CONFIG_INVALID")
            yield descriptor
        except OSError:
            raise ReportStoreError("REPORT_ARTIFACT_UNAVAILABLE") from None
        finally:
            if descriptor is not None:
                os.close(descriptor)

    @staticmethod
    def _read_artifact(directory: int, digest: str) -> bytes | None:
        if not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        try:
            descriptor = os.open(
                f"{digest}.xlsx", os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK,
                dir_fd=directory,
            )
        except FileNotFoundError:
            return None
        with os.fdopen(descriptor, "rb") as handle:
            info = os.fstat(handle.fileno())
            if (
                not stat.S_ISREG(info.st_mode)
                or info.st_mode & 0o077
                or not 0 < info.st_size <= MAX_ARTIFACT_BYTES
            ):
                raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
            content = handle.read(MAX_ARTIFACT_BYTES + 1)
        if len(content) > MAX_ARTIFACT_BYTES or hashlib.sha256(content).hexdigest() != digest:
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        return content

    def _cache_artifact(self, content: bytes, digest: str) -> None:
        if self._artifact_dir is None:
            return
        if not 0 < len(content) <= MAX_ARTIFACT_BYTES:
            raise ReportStoreError("REPORT_SIZE_INVALID")
        if hashlib.sha256(content).hexdigest() != digest:
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        with self._artifact_directory() as directory:
            if self._read_artifact(directory, digest) is not None:
                return
            temporary = f".artifact-{uuid.uuid4().hex}"
            descriptor = os.open(
                temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW,
                0o600, dir_fd=directory,
            )
            try:
                with os.fdopen(descriptor, "wb") as handle:
                    handle.write(content)
                    handle.flush()
                    os.fsync(handle.fileno())
                # Atomic no-overwrite publication also handles concurrent saves.
                try:
                    os.link(
                        temporary, f"{digest}.xlsx", src_dir_fd=directory,
                        dst_dir_fd=directory, follow_symlinks=False,
                    )
                except FileExistsError:
                    pass
            finally:
                os.unlink(temporary, dir_fd=directory)
            os.fsync(directory)
            if self._read_artifact(directory, digest) != content:
                raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")

    @contextmanager
    def _connect(self):
        try:
            with psycopg.connect(
                self._database_url,
                connect_timeout=15,
                sslmode="verify-full",
                row_factory=dict_row,
                prepare_threshold=None,
            ) as connection:
                connection.execute("SET LOCAL statement_timeout = '120s'")
                connection.execute("SET LOCAL lock_timeout = '5s'")
                yield connection
        except psycopg.Error:
            raise ReportStoreError("ONLINE_STORAGE_UNAVAILABLE") from None

    def save(self, result: DraftResult, *, connection=None) -> dict[str, Any]:
        """Verify and save sheets; a supplied transaction owns commit/rollback."""
        evidence = result.evidence
        encoded_payload = _json(result.payload)
        payload_sha = hashlib.sha256(encoded_payload).hexdigest()
        if len(encoded_payload) > MAX_PAYLOAD_BYTES:
            raise ReportStoreError("REPORT_SIZE_INVALID")
        if (
            not re.fullmatch(r"[0-9a-f]{64}", result.input_hash)
            or result.workbook_sha256 != hashlib.sha256(result.xlsx).hexdigest()
            or evidence.get("input_hash") != result.input_hash
            or evidence.get("workbook_sha256") != result.workbook_sha256
            or evidence.get("payload_sha256") != payload_sha
            or evidence.get("state") != "draft"
            or any(
                evidence.get(key, {}).get("status") != "PASS"
                for key in (
                    "payload_validation",
                    "parquet_validation",
                    "workbook_validation",
                )
            )
        ):
            raise ReportStoreError("REPORT_VALIDATION_REQUIRED")
        validation = online_workbook.validate_workbook(result.xlsx, result.payload)
        if validation.get("status") != "PASS" or validation.get("failure_count") != 0:
            raise ReportStoreError("REPORT_WORKBOOK_INVALID")
        snapshots = snapshot_workbook(result.xlsx)
        snapshot_sha = _sha(snapshots)
        self._cache_artifact(result.xlsx, result.workbook_sha256)
        fingerprint = renderer_fingerprint()
        report_id = uuid.uuid4().hex
        transaction = (
            nullcontext(connection) if connection is not None else self._connect()
        )
        with transaction as connection:
            inserted = connection.execute(
                "INSERT INTO psi_preview.reports "
                "(id, as_of, input_hash, workbook_sha256, payload_sha256, "
                "renderer_fingerprint, snapshot_sha256, payload, payload_json, evidence, summary, gates, sheet_count) "
                "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,17) "
                "ON CONFLICT (input_hash, workbook_sha256) DO NOTHING RETURNING id",
                (
                    report_id,
                    result.payload["as_of"],
                    result.input_hash,
                    result.workbook_sha256,
                    payload_sha,
                    fingerprint,
                    snapshot_sha,
                    Jsonb(result.payload),
                    json.dumps(result.payload, ensure_ascii=False, allow_nan=False),
                    Jsonb(evidence),
                    Jsonb(result.payload.get("summary", {})),
                    Jsonb(result.payload.get("gates", [])),
                ),
            ).fetchone()
            if inserted is None:
                existing = connection.execute(
                    "SELECT id, payload_sha256, renderer_fingerprint, snapshot_sha256, evidence "
                    "FROM psi_preview.reports WHERE input_hash=%s AND workbook_sha256=%s",
                    (result.input_hash, result.workbook_sha256),
                ).fetchone()
                if existing is None or any(
                    (
                        existing["payload_sha256"] != payload_sha,
                        existing["renderer_fingerprint"] != fingerprint,
                        existing["snapshot_sha256"] != snapshot_sha,
                        _json(existing["evidence"]) != _json(evidence),
                    )
                ):
                    raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
                report_id = existing["id"]
            else:
                for snapshot in snapshots:
                    connection.execute(
                        "INSERT INTO psi_preview.sheets "
                        "(report_id, ordinal, name, row_count, column_count, populated_rows, cell_count, sha256) "
                        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
                        (
                            report_id,
                            *(
                                snapshot[key]
                                for key in (
                                    "ordinal",
                                    "name",
                                    "row_count",
                                    "column_count",
                                    "populated_rows",
                                    "cell_count",
                                    "sha256",
                                )
                            ),
                        ),
                    )
                with connection.cursor().copy(
                    "COPY psi_preview.sheet_rows (report_id, sheet_ordinal, row_number, cells) FROM STDIN"
                ) as copy:
                    for snapshot in snapshots:
                        for row in snapshot["rows"]:
                            copy.write_row(
                                (
                                    report_id,
                                    snapshot["ordinal"],
                                    row["row_number"],
                                    _json(row["cells"]).decode(),
                                )
                            )
            restored = connection.execute(
                "SELECT payload, payload_json, evidence FROM psi_preview.reports WHERE id=%s",
                (report_id,),
            ).fetchone()
            if (
                _json(restored["payload"]) != encoded_payload
                or restored["payload_json"]
                != json.dumps(result.payload, ensure_ascii=False, allow_nan=False)
                or _json(restored["evidence"]) != _json(evidence)
                or self._snapshots(connection, report_id) != snapshots
            ):
                raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
            row = connection.execute(
                f"SELECT {REPORT_FIELDS} FROM psi_preview.reports WHERE id=%s",
                (report_id,),
            ).fetchone()
            return _metadata(row)

    @staticmethod
    def _snapshots(connection, report_id: str) -> list[dict[str, Any]]:
        sheets = connection.execute(
            f"SELECT {SHEET_FIELDS} FROM psi_preview.sheets WHERE report_id=%s ORDER BY ordinal",
            (report_id,),
        ).fetchall()
        for sheet in sheets:
            sheet["rows"] = connection.execute(
                "SELECT row_number, cells FROM psi_preview.sheet_rows "
                "WHERE report_id=%s AND sheet_ordinal=%s ORDER BY row_number",
                (report_id, sheet["ordinal"]),
            ).fetchall()
        return sheets

    def list(self, limit: int = 20, offset: int = 0) -> list[dict[str, Any]]:
        """List bounded Draft metadata, newest first; never retrieve full payloads."""
        _page(limit, offset, 100)
        with self._connect() as connection:
            return [
                _metadata(row)
                for row in connection.execute(
                    f"SELECT {REPORT_FIELDS} FROM psi_preview.reports "
                    "ORDER BY created_at DESC, id DESC LIMIT %s OFFSET %s",
                    (limit, offset),
                ).fetchall()
            ]

    def get(self, report_id: str) -> dict[str, Any] | None:
        """Read a report's metadata."""
        with self._connect() as connection:
            row = connection.execute(
                f"SELECT {REPORT_FIELDS} FROM psi_preview.reports WHERE id=%s",
                (_report_id(report_id),),
            ).fetchone()
            return _metadata(row) if row else None

    def sheets(self, report_id: str) -> list[dict[str, Any]]:
        """List sheet names, dimensions and verified content hashes."""
        with self._connect() as connection:
            self._require_report(connection, report_id)
            return connection.execute(
                f"SELECT {SHEET_FIELDS} FROM psi_preview.sheets WHERE report_id=%s ORDER BY ordinal",
                (_report_id(report_id),),
            ).fetchall()

    def sheet_rows(
        self, report_id: str, sheet_ordinal: int, limit: int = 100, offset: int = 0
    ) -> list[dict[str, Any]]:
        """Read a bounded page of sparse rows; numbers are Excel's 1-based positions."""
        _page(limit, offset, 500)
        if type(sheet_ordinal) is not int or not 1 <= sheet_ordinal <= 17:
            raise ReportStoreError("REPORT_SHEET_INVALID")
        with self._connect() as connection:
            self._require_report(connection, report_id)
            return connection.execute(
                "SELECT row_number, cells FROM psi_preview.sheet_rows "
                "WHERE report_id=%s AND sheet_ordinal=%s ORDER BY row_number LIMIT %s OFFSET %s",
                (_report_id(report_id), sheet_ordinal, limit, offset),
            ).fetchall()

    @staticmethod
    def _require_report(connection, report_id: str) -> None:
        if (
            connection.execute(
                "SELECT id FROM psi_preview.reports WHERE id=%s",
                (_report_id(report_id),),
            ).fetchone()
            is None
        ):
            raise ReportStoreError("ONLINE_REPORT_NOT_FOUND")

    def payload(self, report_id: str) -> dict[str, Any]:
        """Read the original ordered payload only after verifying both stored forms."""
        with self._connect() as connection:
            row = connection.execute(
                "SELECT payload, payload_json, payload_sha256 FROM psi_preview.reports WHERE id=%s",
                (_report_id(report_id),),
            ).fetchone()
        if row is None:
            raise ReportStoreError("ONLINE_REPORT_NOT_FOUND")
        try:
            payload = json.loads(row["payload_json"])
            valid = (
                isinstance(payload, dict)
                and _sha(payload) == row["payload_sha256"]
                and _json(payload) == _json(row["payload"])
            )
        except (TypeError, ValueError):
            valid = False
        if not valid:
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        return payload

    def download(self, report_id: str) -> bytes:
        """Rebuild the exact validated Draft from a hash-verified online payload."""
        with self._connect() as connection:
            row = connection.execute(
                "SELECT payload, payload_json, payload_sha256, renderer_fingerprint, workbook_sha256, snapshot_sha256 "
                "FROM psi_preview.reports WHERE id=%s",
                (_report_id(report_id),),
            ).fetchone()
        if row is None:
            raise ReportStoreError("ONLINE_REPORT_NOT_FOUND")
        try:
            ordered_payload = json.loads(row["payload_json"])
        except (TypeError, ValueError):
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED") from None
        if _sha(ordered_payload) != row["payload_sha256"] or _json(
            ordered_payload
        ) != _json(row["payload"]):
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        content = None
        if self._artifact_dir is not None:
            with self._artifact_directory() as directory:
                content = self._read_artifact(directory, row["workbook_sha256"])
        if content is None and row["renderer_fingerprint"] != renderer_fingerprint():
            raise ReportStoreError("ONLINE_REPORT_VERSION_UNSUPPORTED")
        try:
            if content is None:
                content = online_workbook.build_workbook(ordered_payload)
            validation = online_workbook.validate_workbook(content, ordered_payload)
            if (
                hashlib.sha256(content).hexdigest() != row["workbook_sha256"]
                or validation.get("status") != "PASS"
                or validation.get("failure_count") != 0
            ):
                raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
            if _sha(snapshot_workbook(content)) != row["snapshot_sha256"]:
                raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED")
        except ReportStoreError:
            raise
        except Exception:
            raise ReportStoreError("ONLINE_REPORT_INTEGRITY_FAILED") from None
        self._cache_artifact(content, row["workbook_sha256"])
        return content
