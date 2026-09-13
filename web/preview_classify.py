"""Identify uploaded exports by the workbook contract, never their filename."""

from __future__ import annotations

import io
import unicodedata
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook

from psi_tool._online_manual_check import load_manual_check

# Sheet names and header positions match the canonical pipeline. Classification
# is routing only: the independent full-data validators still run for each Draft.
SCHEMA_ERROR = "SOURCE_SCHEMA_INVALID"

CONTRACTS = {
    "crm": (
        (
            "Danh sách",
            1,
            (
                ("Số đơn hàng",),
                ("Ngày duyệt",),
                ("Trạng thái phê duyệt",),
                ("Tình trạng",),
            ),
        ),
        (
            "Bảng hàng hóa",
            1,
            (
                ("Số đơn hàng",),
                ("Mã hàng hóa",),
                ("Số lượng",),
                ("Thành tiền sau CK", "Tổng tiền"),
            ),
        ),
    ),
    "product": (
        (
            "Danh sách",
            1,
            (
                ("Mã hàng hóa",),
                ("Tên hàng hóa",),
                ("Nguồn gốc",),
                ("Category",),
                ("Sub Category",),
            ),
        ),
    ),
    "revenue": (
        (
            "SỔ CHI TIẾT BÁN HÀNG",
            4,
            (
                ("Ngày hạch toán",),
                ("Mã hàng",),
                ("Tổng số lượng bán",),
                ("Doanh số bán",),
                ("Chiết khấu",),
                ("Giá trị trả lại",),
                ("Giá trị giảm giá",),
                ("Giá vốn",),
                ("Đơn hàng",),
                ("TK Nợ",),
                ("TK Có",),
            ),
        ),
    ),
    "inventory": (
        (
            "TỔNG HỢP TỒN KHO",
            4,
            (
                ("Tên kho",),
                ("Mã kho",),
                ("Mã hàng",),
                ("Cuối kỳ",),
            ),
        ),
    ),
}


def _normalize(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join(
        "".join(c for c in text if not unicodedata.combining(c)).casefold().split()
    )


def _headers(book: Workbook, sheet: str, row: int) -> tuple:
    if sheet not in book.sheetnames:
        return ()
    return next(
        book[sheet].iter_rows(
            min_row=row,
            max_row=row,
            max_col=256,
            values_only=True,
        ),
        (),
    )


def classify_files(files: list[tuple[str, bytes]]) -> list[dict]:
    """Return one result per input, retaining ambiguity and duplicates for UI."""
    results = []
    for index, (filename, content) in enumerate(files):
        candidates = []
        reason = "Không khớp cấu trúc của bốn nguồn kỳ báo cáo."
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                for role, sheets in CONTRACTS.items():
                    matched = True
                    for sheet, row, requirements in sheets:
                        headers = {_normalize(x) for x in _headers(book, sheet, row)}
                        if not all(
                            any(_normalize(alias) in headers for alias in alternatives)
                            for alternatives in requirements
                        ):
                            matched = False
                            break
                    if matched:
                        candidates.append(role)
            finally:
                book.close()
            if len(candidates) == 1:
                reason = "Đã nhận diện từ tên sheet và các cột dữ liệu."
            elif candidates:
                reason = "File khớp nhiều loại nguồn; cần chọn loại phù hợp."
        except (ValueError, KeyError, OSError, EOFError, BadZipFile, ParseError):
            reason = "Không đọc được cấu trúc Excel; kiểm tra lại file."
        results.append(
            {
                "index": index,
                "filename": Path(filename.replace("\\", "/")).name,
                "role": candidates[0] if len(candidates) == 1 else None,
                "candidates": candidates,
                "reason": reason,
            }
        )
    return results


def validate_reusable(role: str, filename: str, content: bytes) -> None:
    """Reject obvious wrong-role files without inventing an approval decision."""
    del filename  # Names are display metadata, not evidence of a file's role.
    if role == "manual_check":
        try:
            load_manual_check(io.BytesIO(content))
        except (ValueError, KeyError) as exc:
            raise ValueError(SCHEMA_ERROR) from exc
        return
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if role == "purchase":
            headers = _headers(book, "LDL", 4)
            expected = {
                18: "MÃ MISA",
                19: "TÊN HÀNG HÓA",
                20: "SL",
                76: "Giá nhập kho (chưa VAT)",
                77: "Giá trị nhập kho (theo SL)",
            }
            valid = all(
                len(headers) > i and _normalize(headers[i]) == _normalize(label)
                for i, label in expected.items()
            )
        elif role == "target":
            headers = {_normalize(x) for x in _headers(book, "Target", 2)}
            valid = {"brand", "brand code"} <= headers and any(
                "target" in name for name in headers
            )
        else:
            valid = False
        if not valid:
            raise ValueError(SCHEMA_ERROR)
    finally:
        book.close()
