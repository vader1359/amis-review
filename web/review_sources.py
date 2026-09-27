"""Immutable private source bundles and governed accounting exclusions."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import shutil
import stat
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from psi_tool._online_manual_check import (
    ORDER_EXCLUSION_HEADERS,
    PREORDER_HEADERS,
    load_manual_check,
    make_order_exclusion_fingerprint,
    make_preorder_fingerprint,
)
from psi_tool.online_pipeline import MAX_SOURCE_BYTES, SOURCE_LABELS, SourceSnapshot


class SourceBundleError(ValueError):
    def __init__(self, code):
        self.code = code
        super().__init__(code)


def _digest(content):
    return hashlib.sha256(content).hexdigest()


def _json(value):
    return json.dumps(
        value, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    ).encode()


def _check_path(path):
    if any(item.is_symlink() for item in [path, *path.parents]):
        raise SourceBundleError("SOURCE_BUNDLE_UNSAFE_PATH")


def _read(path, limit):
    _check_path(path)
    fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
    with os.fdopen(fd, "rb") as stream:
        info = os.fstat(stream.fileno())
        if not stat.S_ISREG(info.st_mode) or info.st_size > limit:
            raise SourceBundleError("SOURCE_BUNDLE_INVALID")
        content = stream.read(limit + 1)
    if len(content) > limit:
        raise SourceBundleError("SOURCE_BUNDLE_INVALID")
    return content


def _write(path, content):
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600)
    with os.fdopen(fd, "wb") as stream:
        stream.write(content)
        stream.flush()
        os.fsync(stream.fileno())


def _filename(value):
    return (
        isinstance(value, str)
        and value.lower().endswith(".xlsx")
        and not any(c in value for c in ("/", "\\", "\x00"))
    )


class SourceBundleStore:
    def __init__(self, directory: Path):
        self.directory = Path(directory).absolute()
        _check_path(self.directory)
        self.directory.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.directory.chmod(0o700)

    def _path(self, report_id):
        if not isinstance(report_id, str) or not re.fullmatch(
            r"[a-f0-9]{32}", report_id
        ):
            raise SourceBundleError("SOURCE_BUNDLE_ID_INVALID")
        path = self.directory / report_id
        _check_path(path)
        return path

    def save(self, report_id, snapshots, baseline):
        destination = self._path(report_id)
        if set(snapshots) != set(SOURCE_LABELS):
            raise SourceBundleError("SOURCE_SET_INVALID")
        selected = {**snapshots, "prior_psi": baseline}
        records = {}
        for role, snapshot in selected.items():
            if (
                not isinstance(snapshot, SourceSnapshot)
                or not isinstance(snapshot.content, bytes)
                or not 0 < len(snapshot.content) <= MAX_SOURCE_BYTES
                or _digest(snapshot.content) != snapshot.sha256
                or not _filename(snapshot.filename)
            ):
                raise SourceBundleError("SOURCE_SNAPSHOT_INVALID")
            records[role] = {
                "filename": snapshot.filename,
                "sha256": snapshot.sha256,
                "size": len(snapshot.content),
            }
        manifest = {"version": 1, "report_id": report_id, "sources": records}
        envelope = _json({"manifest": manifest, "sha256": _digest(_json(manifest))})
        if destination.exists():
            existing, prior = self.load(report_id)
            if existing != snapshots or prior != baseline:
                raise SourceBundleError("SOURCE_BUNDLE_IMMUTABLE")
            return
        temporary = Path(tempfile.mkdtemp(prefix=".bundle-", dir=self.directory))
        try:
            for snapshot in selected.values():
                blob = temporary / (snapshot.sha256 + ".xlsx")
                if not blob.exists():
                    _write(blob, snapshot.content)
            _write(temporary / "manifest.json", envelope)
            try:
                os.rename(temporary, destination)
            except OSError:
                existing, prior = self.load(report_id)
                if existing != snapshots or prior != baseline:
                    raise SourceBundleError("SOURCE_BUNDLE_IMMUTABLE") from None
        finally:
            if temporary.exists():
                shutil.rmtree(temporary)

    def load(self, report_id):
        directory = self._path(report_id)
        try:
            envelope = json.loads(_read(directory / "manifest.json", 32768))
            manifest = envelope["manifest"]
            if (
                _digest(_json(manifest)) != envelope["sha256"]
                or manifest["version"] != 1
                or manifest["report_id"] != report_id
                or set(manifest["sources"]) != set(SOURCE_LABELS) | {"prior_psi"}
            ):
                raise SourceBundleError("SOURCE_BUNDLE_INVALID")
            snapshots = {}
            for role, record in manifest["sources"].items():
                checksum = record["sha256"]
                if not re.fullmatch(r"[a-f0-9]{64}", checksum) or not _filename(
                    record["filename"]
                ):
                    raise SourceBundleError("SOURCE_BUNDLE_INVALID")
                content = _read(directory / (checksum + ".xlsx"), MAX_SOURCE_BYTES)
                if (
                    not content
                    or len(content) != record["size"]
                    or _digest(content) != checksum
                ):
                    raise SourceBundleError("SOURCE_BUNDLE_CHECKSUM_MISMATCH")
                snapshots[role] = SourceSnapshot(record["filename"], content, checksum)
            baseline = snapshots.pop("prior_psi")
            return snapshots, baseline
        except FileNotFoundError:
            raise SourceBundleError("SOURCE_BUNDLE_NOT_FOUND") from None
        except (OSError, KeyError, TypeError, json.JSONDecodeError):
            raise SourceBundleError("SOURCE_BUNDLE_INVALID") from None

    def available(self, report_id):
        try:
            self.load(report_id)
            return True
        except SourceBundleError:
            return False


def _append_record(sheet, headers, record):
    for row in sheet.iter_rows(min_row=1, max_row=10):
        columns = {
            str(cell.value).strip(): cell.column
            for cell in row
            if cell.value is not None
        }
        if set(headers) <= set(columns):
            target = sheet.max_row + 1
            for key, value in record.items():
                cell = sheet.cell(target, columns[key], value)
                if isinstance(value, str):
                    cell.data_type = "s"
            return
    raise SourceBundleError("MANUAL_CHECK_HEADERS_INVALID")


def apply_exclusions(snapshots, proposals, as_of: date, author: str):
    if (
        set(snapshots) != set(SOURCE_LABELS)
        or type(as_of) is not date
        or not author.strip()
    ):
        raise SourceBundleError("REVIEW_APPROVAL_INVALID")
    original = snapshots["manual_check"]
    if _digest(original.content) != original.sha256:
        raise SourceBundleError("SOURCE_CHECKSUM_MISMATCH")
    registry = load_manual_check(original.content)
    workbook = load_workbook(io.BytesIO(original.content))
    changed = False
    seen = set()
    try:
        existing = {
            rule.exclusion_id: rule
            for rule in (*registry.preorder_exclusions, *registry.order_exclusions)
        }
        for proposal in proposals:
            identifier = str(proposal.get("id", "")).strip()
            order = str(proposal.get("order_id", "")).strip().upper()
            sku = str(proposal.get("sku", "")).strip().upper()
            note = str(proposal.get("note", "")).strip()
            proposer = str(proposal.get("author", "")).strip()
            kind = proposal.get("action")
            if (
                not identifier
                or len(identifier) > 128
                or not order
                or not note
                or not proposer
            ):
                raise SourceBundleError("REVIEW_PROPOSAL_INVALID")
            if identifier in seen:
                raise SourceBundleError("REVIEW_PROPOSAL_DUPLICATE")
            seen.add(identifier)
            exclusion_id = "ONLINE-" + _digest(identifier.encode())[:32]
            common = {
                "Exclusion ID": exclusion_id,
                "Status": "APPROVED",
                "Order ID": order,
                "Effective From": as_of.isoformat(),
                "Approved By": author.strip(),
                "Approved Date": as_of.isoformat(),
                "Disposition": "PERMANENT / DONE",
                "Evidence": "Online accounting review proposal " + identifier,
                "Notes": "Proposed by: " + proposer + "; " + note,
            }
            if kind == "exclude_preorder" and sku:
                canonical = registry.map_sku(sku, as_of=as_of)
                fingerprint = make_preorder_fingerprint(
                    action="EXCLUDE FROM PREORDER",
                    order_id=order,
                    canonical_sku=canonical,
                )
                record = {
                    **common,
                    "Fingerprint": fingerprint,
                    "Action": "EXCLUDE FROM PREORDER",
                    "Raw SKU": sku,
                    "Canonical SKU": canonical,
                    "KT Note": note,
                    "Treatment": "EXCLUDE FROM PREORDER",
                }
                sheet, headers = "Preorder Exclusions", PREORDER_HEADERS
            elif kind == "exclude_order":
                fingerprint = make_order_exclusion_fingerprint(
                    action="EXCLUDE ORDER FROM PSI",
                    scope="ALL PSI BUSINESS SHEETS",
                    order_id=order,
                )
                record = {
                    **common,
                    "Fingerprint": fingerprint,
                    "Action": "EXCLUDE ORDER FROM PSI",
                    "Scope": "ALL PSI BUSINESS SHEETS",
                    "Reason": note,
                    "Treatment": "EXCLUDE ORDER FROM PSI",
                }
                sheet, headers = "Order Exclusions", ORDER_EXCLUSION_HEADERS
            else:
                raise SourceBundleError("REVIEW_ACTION_INVALID")
            if exclusion_id in existing:
                rule = existing[exclusion_id]
                if (
                    rule.fingerprint != fingerprint
                    or rule.evidence != common["Evidence"]
                    or rule.notes != common["Notes"]
                    or rule.approved_by != author.strip()
                    or rule.effective_from != as_of
                    or rule.approved_date != as_of
                    or rule.effective_to is not None
                    or rule.disposition != "PERMANENT / DONE"
                    or rule.status != "APPROVED"
                ):
                    raise SourceBundleError("REVIEW_PROPOSAL_CONFLICT")
                continue
            _append_record(workbook[sheet], headers, record)
            changed = True
        if not changed:
            return dict(snapshots)
        output = io.BytesIO()
        workbook.properties.created = datetime(1980, 1, 1, tzinfo=timezone.utc)
        workbook.save(output)
        normalized = io.BytesIO()
        with ZipFile(output) as source, ZipFile(normalized, "w") as target:
            for entry in source.infolist():
                entry.date_time = (1980, 1, 1, 0, 0, 0)
                data = source.read(entry.filename)
                if entry.filename == "docProps/core.xml":
                    data = re.sub(
                        rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>1980-01-01T00:00:00Z\g<2>",
                        data,
                    )
                target.writestr(entry, data)
        content = normalized.getvalue()
        load_manual_check(content)
        return {
            **snapshots,
            "manual_check": SourceSnapshot.from_bytes(original.filename, content),
        }
    finally:
        workbook.close()


def merge_applied_exclusions(snapshots, decisions, as_of: date):
    """Replay applied decisions using their original effective date and approver.

    Caller supplies only applied decisions from the authoritative review store.
    """
    if type(as_of) is not date:
        raise SourceBundleError("REVIEW_APPROVAL_INVALID")
    groups = {}
    for decision in decisions:
        try:
            effective = date.fromisoformat(decision["effective_from"])
            author = decision["approved_by"]
        except (KeyError, TypeError, ValueError):
            raise SourceBundleError("REVIEW_APPROVAL_INVALID") from None
        if not isinstance(author, str) or not author.strip():
            raise SourceBundleError("REVIEW_APPROVAL_INVALID")
        if effective <= as_of:
            groups.setdefault((effective, author.strip()), []).append(decision)
    merged = dict(snapshots)
    for (effective, author), proposals in sorted(groups.items()):
        merged = apply_exclusions(merged, proposals, effective, author)
    return merged
